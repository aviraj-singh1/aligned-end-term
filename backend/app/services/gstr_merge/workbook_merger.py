import copy
import os
from pathlib import Path
import shutil
import tempfile
from typing import Any, Dict, List
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from backend.app.services.gstr_merge.gst_schema import CATEGORY_SPECS, normalize_label
from backend.app.services.gstr_merge.header_utils import annualize_metadata_value
from backend.app.services.gstr_merge.source_manifest import (
    SourceManifestBuilder,
    SourceValidationError,
    iter_spooled_rows,
    validate_excel_row_limits,
)


class MergeValidationError(ValueError):
    pass


class WorkbookMerger:
    """Stream validated rows into a write-only workbook."""

    def __init__(self, temp_paths: Dict[str, str], metadata: Dict[str, Any], output_path: str,
                 source_manifest: Dict[str, Any] = None, mode: str = "quarterly"):
        self.temp_paths = temp_paths
        self.metadata = metadata
        self.output_path = output_path
        self.manifest = source_manifest
        self.mode = mode
        self.duplicate_warnings: List[Dict[str, Any]] = []

    @staticmethod
    def _consistent_value(items, key):
        values = {item.get(key) for item in items if item.get(key)}
        return next(iter(values)) if len(values) == 1 else None

    @staticmethod
    def _period_label(period):
        if not period:
            return "Not verifiable"
        months = {
            "M01": "January", "M02": "February", "M03": "March", "M04": "April",
            "M05": "May", "M06": "June", "M07": "July", "M08": "August",
            "M09": "September", "M10": "October", "M11": "November", "M12": "December",
        }
        return months.get(period["code"], period["code"])

    def _write_annual_summary(self, output_wb):
        ws = output_wb.create_sheet("Annual Summary")
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A12"
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 38
        ws.column_dimensions["C"].width = 24
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 22
        items = list(self.metadata.values())
        expected_codes = (
            ["Q1", "Q2", "Q3", "Q4"] if self.mode == "quarterly"
            else [f"M{month:02d}" for month in range(4, 13)] + [f"M{month:02d}" for month in range(1, 4)]
        )
        periods = [item.get("period") for item in items]
        actual_codes = {period["code"] for period in periods if period}
        included = [self._period_label({"code": code}) for code in expected_codes if code in actual_codes]
        missing = [self._period_label({"code": code}) for code in expected_codes if code not in actual_codes]

        title = WriteOnlyCell(ws, value="Annual GSTR Merge Summary")
        title.font = Font(size=16, bold=True, color="FFFFFF")
        title.fill = PatternFill("solid", fgColor="1D4ED8")
        title.alignment = Alignment(vertical="center")
        ws.append([title, "", "", "", ""])
        ws.append(["Integrity status", "Verified after merge", "Mode", self.mode.title()])
        ws.append(["Financial year", self._consistent_value(items, "financial_year") or "Not verifiable",
                   "Taxpayer GSTIN", self._consistent_value(items, "taxpayer_gstin") or "Not verifiable"])
        ws.append(["Legal name", self._consistent_value(items, "legal_name") or "Not verifiable",
                   "Trade name", self._consistent_value(items, "trade_name") or "Not verifiable"])
        ws.append(["Included periods", ", ".join(included) or "Not verifiable",
                   "Missing periods", ", ".join(missing) or "None"])
        ws.append(["Source files", len(items), "Merged detail rows",
                   self.manifest.get(
                       "annual_summary_row_count",
                       sum(category["row_count"] for category in self.manifest.get("categories", {}).values()),
                   )])
        ws.append([])
        ws.append(["Notes"])
        has_omitted_summaries = any(
            info.get("static_policy") == "ignored_summary"
            for item in items for info in item.get("sheets", {}).values()
        )
        has_read_me = any(
            normalize_label(info.get("sheet_name")) == "read me"
            for item in items for info in item.get("sheets", {}).values()
        )
        if has_omitted_summaries:
            ws.append(["Period-level ITC summary sheets are not arithmetically combined; annual figures are verified from merged detail rows."])
        if has_read_me:
            ws.append(["The retained 'Read me' sheet is source guidance from the first workbook, not an annual-period declaration."])
        ws.append([])

        headers = ["Source file", "Detected tax period", "Metadata source", "Detail rows included", "Generated on"]
        header_cells = []
        for value in headers:
            cell = WriteOnlyCell(ws, value=value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="475569")
            header_cells.append(cell)
        ws.append(header_cells)
        for source_id, item in self.metadata.items():
            period = item.get("period")
            row_count = sum(
                category.get("source_counts", {}).get(source_id, 0)
                for category in self.manifest.get("categories", {}).values()
            )
            ws.append([
                item["filename"], self._period_label(period),
                period.get("source", "Unavailable").title() if period else "Unavailable",
                row_count,
                item.get("generation_date") or "Not available",
            ])

    @staticmethod
    def _styled_cell(output_ws, source_cell, value=None, full_style=True):
        target = WriteOnlyCell(output_ws, value=source_cell.value if value is None else value)
        if source_cell.has_style:
            if full_style:
                target.font = copy.copy(source_cell.font)
                target.fill = copy.copy(source_cell.fill)
                target.border = copy.copy(source_cell.border)
                target.alignment = copy.copy(source_cell.alignment)
                target.protection = copy.copy(source_cell.protection)
            target.number_format = source_cell.number_format
        return target

    @staticmethod
    def _copy_dimensions(source_ws, output_ws, max_row=None):
        for key, dimension in source_ws.column_dimensions.items():
            target = output_ws.column_dimensions[key]
            target.width = dimension.width
            target.hidden = dimension.hidden
            target.bestFit = dimension.bestFit
            target.outlineLevel = dimension.outlineLevel
        for index, dimension in source_ws.row_dimensions.items():
            if max_row is not None and index > max_row:
                continue
            target = output_ws.row_dimensions[index]
            target.height = dimension.height
            target.hidden = dimension.hidden
            target.outlineLevel = dimension.outlineLevel
        output_ws.freeze_panes = source_ws.freeze_panes
        output_ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
        output_ws.sheet_state = source_ws.sheet_state

    def _write_static_sheet(self, output_wb, source_ws, merged_ranges):
        output_ws = output_wb.create_sheet(source_ws.title)
        self._copy_dimensions(source_ws, output_ws)
        for row in source_ws.iter_rows():
            output_row = []
            first_value = row[0].value if row else None
            for cell in row:
                value = cell.value
                if cell.column == 2 and normalize_label(first_value) == "months":
                    value = annualize_metadata_value(first_value, value)
                output_row.append(self._styled_cell(output_ws, cell, value=value, full_style=True))
            output_ws.append(output_row)
        ranges = [str(item) for item in source_ws.merged_cells.ranges]
        if ranges:
            merged_ranges[output_ws.title] = ranges

    def _write_business_template(self, output_wb, source_ws, info, category_manifest, merged_ranges):
        output_ws = output_wb.create_sheet(category_manifest["sheet_name"])
        header_row = info["header_row_idx"]
        self._copy_dimensions(source_ws, output_ws, max_row=header_row)
        for row in source_ws.iter_rows(min_row=1, max_row=header_row):
            output_ws.append([self._styled_cell(output_ws, cell, full_style=True) for cell in row])

        formats = ["General"] * category_manifest["output_width"]
        if source_ws.max_row > header_row:
            sample = source_ws[header_row + 1]
            source_map = info["column_map"]
            for key, target_index in category_manifest["column_map"].items():
                source_index = source_map.get(key)
                if source_index is not None and source_index < len(sample):
                    formats[target_index] = sample[source_index].number_format
        category_manifest["number_formats"] = tuple(formats)
        ranges = [str(item) for item in source_ws.merged_cells.ranges if item.max_row <= header_row]
        if ranges:
            merged_ranges[output_ws.title] = ranges
        return output_ws

    @staticmethod
    def _inject_merged_ranges(path: str, ranges_by_title: Dict[str, List[str]]):
        if not ranges_by_title:
            return
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with zipfile.ZipFile(path, "r") as source:
            workbook = ET.fromstring(source.read("xl/workbook.xml"))
            relationships = ET.fromstring(source.read("xl/_rels/workbook.xml.rels"))
            targets = {item.attrib["Id"]: item.attrib["Target"] for item in relationships}
            sheet_paths = {}
            for sheet in workbook.find(f"{{{namespace}}}sheets"):
                rel_id = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                target = targets[rel_id].lstrip("/")
                sheet_paths[sheet.attrib["name"]] = target if target.startswith("xl/") else f"xl/{target}"
            patches = {sheet_paths[title]: ranges for title, ranges in ranges_by_title.items() if title in sheet_paths}
            handle = tempfile.NamedTemporaryFile(
                prefix="gstr-merged-", suffix=".xlsx", dir=str(Path(path).parent), delete=False
            )
            patched_path = handle.name
            handle.close()
            try:
                with zipfile.ZipFile(patched_path, "w") as target_zip:
                    for info in source.infolist():
                        ranges = patches.get(info.filename)
                        if not ranges:
                            with source.open(info) as reader, target_zip.open(info, "w") as writer:
                                shutil.copyfileobj(reader, writer, length=1024 * 1024)
                            continue
                        merge_xml = (
                            f'<mergeCells xmlns="{namespace}" count="{len(ranges)}">'
                            + "".join(f'<mergeCell ref="{reference}"/>' for reference in ranges)
                            + "</mergeCells>"
                        ).encode("utf-8")
                        marker = b"</sheetData>"
                        with source.open(info) as reader, target_zip.open(info, "w") as writer:
                            pending = b""
                            inserted = False
                            while chunk := reader.read(1024 * 1024):
                                pending += chunk
                                position = pending.find(marker)
                                if position >= 0 and not inserted:
                                    end = position + len(marker)
                                    writer.write(pending[:end])
                                    writer.write(merge_xml)
                                    pending = pending[end:]
                                    inserted = True
                                elif len(pending) > len(marker):
                                    writer.write(pending[:-(len(marker) - 1)])
                                    pending = pending[-(len(marker) - 1):]
                            writer.write(pending)
                            if not inserted:
                                raise MergeValidationError("Could not preserve template merged cells safely.")
                source.close()
                os.replace(patched_path, path)
            finally:
                if os.path.exists(patched_path):
                    os.remove(patched_path)

    def merge(self):
        source_ids = list(self.temp_paths)
        if not source_ids:
            raise MergeValidationError("No input workbooks were supplied.")
        if self.manifest is None:
            builder = SourceManifestBuilder(self.temp_paths, self.metadata)
            try:
                self.manifest = builder.build()
            except SourceValidationError as exc:
                raise MergeValidationError(str(exc)) from exc
            self.duplicate_warnings = builder.duplicate_warnings
        limit_errors = validate_excel_row_limits(self.manifest)
        if limit_errors:
            raise MergeValidationError("\n".join(limit_errors))

        output_wb = openpyxl.Workbook(write_only=True)
        emitted_static = set()
        business_sheets = {}
        merged_ranges = {}
        partial_path = f"{self.output_path}.partial"
        try:
            self._write_annual_summary(output_wb)
            for source_id in source_ids:
                required_categories = {
                    category for category, item in self.manifest["categories"].items()
                    if item["template_source_id"] == source_id
                }
                static_infos = [
                    info for info in self.metadata[source_id]["sheets"].values()
                    if not info.get("category")
                    and info.get("static_policy") != "ignored_summary"
                    and normalize_label(info["sheet_name"]) != "annual summary"
                    and normalize_label(info["sheet_name"]) not in emitted_static
                ]
                if not required_categories and not static_infos:
                    continue
                source_wb = openpyxl.load_workbook(
                    self.temp_paths[source_id], data_only=False, read_only=False, keep_links=False
                )
                try:
                    for info in static_infos:
                        self._write_static_sheet(output_wb, source_wb[info["sheet_name"]], merged_ranges)
                        emitted_static.add(normalize_label(info["sheet_name"]))
                    for category in required_categories:
                        category_manifest = self.manifest["categories"][category]
                        info = self.metadata[source_id]["categories"][category]
                        business_sheets[category] = self._write_business_template(
                            output_wb, source_wb[info["sheet_name"]], info, category_manifest, merged_ranges
                        )
                finally:
                    source_wb.close()

            for spec in CATEGORY_SPECS:
                category_manifest = self.manifest["categories"].get(spec.key)
                if not category_manifest:
                    continue
                output_ws = business_sheets[spec.key]
                formats = category_manifest["number_formats"]
                styled_columns = tuple(index for index, value in enumerate(formats) if value != "General")
                for _, _, _, _, values in iter_spooled_rows(category_manifest):
                    if not styled_columns:
                        output_ws.append(values)
                        continue
                    row = []
                    for index, value in enumerate(values):
                        cell = WriteOnlyCell(output_ws, value=value)
                        if formats[index] != "General":
                            cell.number_format = formats[index]
                        row.append(cell)
                    output_ws.append(row)

            output_wb.save(partial_path)
            self._inject_merged_ranges(partial_path, merged_ranges)
            os.replace(partial_path, self.output_path)
        finally:
            output_wb.close()
            if os.path.exists(partial_path):
                os.remove(partial_path)
            for category_manifest in self.manifest.get("categories", {}).values():
                spool_path = category_manifest.get("spool_path")
                if spool_path and os.path.exists(spool_path):
                    os.remove(spool_path)
                category_manifest["spool_path"] = None
        return self.manifest
