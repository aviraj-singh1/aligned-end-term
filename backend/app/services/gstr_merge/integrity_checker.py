import decimal
from collections import Counter
from typing import Any, Dict, List

import openpyxl

from backend.app.services.gstr_merge.gst_schema import CATEGORY_BY_KEY, normalize_label
from backend.app.services.gstr_merge.source_manifest import typed_row_token
from backend.app.services.gstr_merge.workbook_inspector import WorkbookInspector


class IntegrityChecker:
    def __init__(self, manifest: Dict[str, Any], merged_path: str, metadata: Dict[str, Any] = None):
        self.manifest = manifest
        self.merged_path = merged_path
        self.metadata = metadata or {}
        self.final_summary: Dict[str, Any] = {}

    @staticmethod
    def _meaningful(value):
        return value is not None and (not isinstance(value, str) or bool(value.strip()))

    @staticmethod
    def _decimal(value, category: str, row: int, field: str, errors: List[str]):
        if value is None or (isinstance(value, str) and not value.strip()):
            return decimal.Decimal("0")
        if (field.startswith("itc_reduction_") and isinstance(value, str)
                and normalize_label(value) in {"na", "n a", "not applicable"}):
            return decimal.Decimal("0")
        try:
            return decimal.Decimal(str(value))
        except (decimal.InvalidOperation, ValueError):
            errors.append(f"{category} row {row}: invalid numeric value in {field}.")
            return decimal.Decimal("0")

    @staticmethod
    def _typed_hash(values):
        return typed_row_token(tuple(values))

    def verify(self) -> List[str]:
        errors: List[str] = []
        try:
            wb = openpyxl.load_workbook(
                self.merged_path, data_only=False, read_only=True, keep_links=False
            )
        except Exception:
            return ["Merged workbook cannot be reopened because its package is invalid."]

        total_rows = 0
        rows_by_category = {}
        expected_titles = {item["sheet_name"] for item in self.manifest.get("categories", {}).values()}
        known_titles = {item.output_name for item in CATEGORY_BY_KEY.values()}
        for title in wb.sheetnames:
            if title in known_titles and title not in expected_titles:
                errors.append(f"Unexpected supported category {title} exists in the output workbook.")
        for category, expected in self.manifest.get("categories", {}).items():
            label = CATEGORY_BY_KEY[category].output_name
            sheet_name = expected["sheet_name"]
            if sheet_name not in wb.sheetnames:
                errors.append(f"Expected category {label} is missing from the output workbook.")
                continue
            ws = wb[sheet_name]
            width = expected["output_width"]
            detected_header, header_error = WorkbookInspector._detect_header(
                ws, CATEGORY_BY_KEY[category]
            )
            if header_error:
                errors.append(f"{label}: output header cannot be verified ({header_error}).")
                continue
            actual_mapping = detected_header["column_map"]
            if detected_header["header_row_idx"] != expected["header_row_idx"]:
                errors.append(f"{label}: output header row changed.")
            for key, column in expected["column_map"].items():
                if actual_mapping.get(key) != column:
                    errors.append(f"{label}: canonical header mapping changed for column {column + 1}.")
            hashes, row_count = Counter(), 0
            totals = {field: decimal.Decimal("0") for field in expected["totals"]}
            for row_number, row in enumerate(
                ws.iter_rows(min_row=expected["header_row_idx"] + 1, max_col=width, values_only=True),
                expected["header_row_idx"] + 1,
            ):
                if not any(self._meaningful(value) for value in row):
                    continue
                row_count += 1
                hashes[self._typed_hash(tuple(row))] += 1
                for field in totals:
                    totals[field] += self._decimal(row[expected["column_map"][field]], label, row_number, field, errors)
            if row_count != expected["row_count"]:
                errors.append(f"{label}: expected {expected['row_count']} rows, found {row_count}.")
            expected_hashes = expected.get("row_hash_counts")
            if expected_hashes is None:
                legacy = [self._typed_hash(tuple(record["values"])) for record in expected.get("rows", [])]
                expected_hashes = Counter(legacy or expected.get("row_hashes", []))
            if hashes != Counter(expected_hashes):
                errors.append(f"{label}: row multiplicity/hash verification failed.")
            for field, expected_total in expected["totals"].items():
                if totals[field] != expected_total:
                    errors.append(f"{label}: {field} total mismatch; expected {expected_total}, found {totals[field]}.")
            rows_by_category[label] = row_count
            total_rows += row_count
        wb.close()
        self.final_summary = {"categories_merged": list(rows_by_category), "rows_merged_per_category": rows_by_category,
                              "total_rows_merged": total_rows, "integrity_status": "passed" if not errors else "failed"}
        return errors

    def get_final_summary(self):
        return self.final_summary
