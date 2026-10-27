"""
File Parser for GST Reconciliation Tool.

Top-level module that:
1. Accepts .xls, .xlsx or .json files
2. Reads all sheets / JSON structures
3. Delegates to sheet_detector for classification
4. Delegates to column_detector for mapping
5. Parses each data row into a ParsedInvoiceRecord using normalizer
6. Returns a FileInspectionResult

This is the single entry point for Phase 1 file ingestion.
"""

from __future__ import annotations

import copy
import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import xlrd

from backend.app.models.records import (
    ColumnMapping,
    FileInspectionResult,
    FinancialParseIssue,
    InvoiceNumberInfo,
    ParsedInvoiceRecord,
    PersistedSheetRows,
    RecordSource,
    SheetCategory,
    SheetDetectionResult,
    VoucherType,
    normalize_note_direction,
)
from backend.app.services.column_detector import build_column_index_map, detect_columns
from backend.app.services.header_detector import detect_header_row
from backend.app.services.header_detector import _normalize_header_text
from backend.app.services.normalizer import (
    normalize_gstin,
    normalize_invoice_number,
    normalize_supplier_name,
    parse_date,
    parse_money,
    parse_rate,
)
from backend.app.services.sheet_detector import detect_sheet


# ---------------------------------------------------------------------------
# Source classification
# ---------------------------------------------------------------------------

_TALLY_CATEGORIES = {
    SheetCategory.PURCHASE_REGISTER,
    SheetCategory.DEBIT_NOTE_REGISTER,
}

_GST_CATEGORIES = {
    SheetCategory.B2B,
    SheetCategory.B2BA,
    SheetCategory.CDNR,
    SheetCategory.CDNRA,
}


def _apply_amendment_column_map(
    headers: List[Any], col_map: Dict[str, int], category: SheetCategory,
) -> None:
    """Force exact portal amendment headers past generic contains matching."""
    if category not in (SheetCategory.B2BA, SheetCategory.CDNRA):
        return
    normalized = {_normalize_header_text(header): index for index, header in enumerate(headers)}
    aliases = {
        "amendment_original_number": (
            "originalinvoiceno.", "originalinvoiceno", "originalinvoicenumber",
            "originalnoteno", "originalnoteno.", "originalnotenumber",
        ),
        "amendment_original_date": ("originalinvoicedate", "originalnotedate"),
        "revised_invoice_number": ("revisedinvoiceno.", "revisedinvoiceno", "revisedinvoicenumber"),
        "revised_invoice_date": ("revisedinvoicedate",),
        "revised_note_number": ("revisednoteno.", "revisednoteno", "revisednotenumber"),
        "revised_note_date": ("revisednotedate",),
        "amendment_date": ("amendment,ifany", "amendmentdate"),
        "amendment_period": ("amendmentperiod",),
    }
    for semantic, names in aliases.items():
        for name in names:
            if name in normalized:
                col_map[semantic] = normalized[name]
                break


def _determine_source(category: SheetCategory) -> RecordSource:
    """Determine record source from sheet category."""
    if category in _GST_CATEGORIES:
        return RecordSource.GST
    return RecordSource.TALLY


def _determine_voucher_type(raw: Any, category: SheetCategory) -> Tuple[VoucherType, str]:
    """Determine the voucher type from raw value and sheet category."""
    original = str(raw).strip() if raw else ""

    if category == SheetCategory.DEBIT_NOTE_REGISTER:
        return (VoucherType.DEBIT_NOTE, original)

    if category in (SheetCategory.CDNR, SheetCategory.CDNRA):
        return (normalize_note_direction(raw), original)

    if original.lower() in ("purchase", "purchases"):
        return (VoucherType.PURCHASE, original)
    if "debit" in original.lower():
        return (VoucherType.DEBIT_NOTE, original)
    if "credit" in original.lower():
        return (VoucherType.CREDIT_NOTE, original)

    return (VoucherType.UNKNOWN, original)


# ---------------------------------------------------------------------------
# Row → Record conversion
# ---------------------------------------------------------------------------

def _row_to_record(
    row_values: List[Any],
    headers: List[Any],
    col_map: Dict[str, int],
    category: SheetCategory,
    sheet_name: str,
    row_number: int,
) -> Optional[ParsedInvoiceRecord]:
    """
    Convert a single data row into a ParsedInvoiceRecord.
    
    Returns None if the row is entirely empty or lacks essential data.
    """

    def _get(field: str) -> Any:
        idx = col_map.get(field)
        if idx is None or idx >= len(row_values):
            return None
        return row_values[idx]

    def _money(field: str, raw: Any = None) -> Decimal:
        value = _get(field) if raw is None else raw
        try:
            return parse_money(value)
        except (ValueError, ArithmeticError):
            record.financial_parse_issues.append(FinancialParseIssue(
                field_name=field,
                raw_value=str(value),
                reason=f"{field} contains an invalid monetary value: {value!r}",
            ))
            return Decimal("0.00")

    # Skip completely empty rows
    non_empty = [v for v in row_values if v is not None and str(v).strip()]
    if not non_empty:
        return None

    # Build raw data dict for auditability
    raw_data = {}
    for i, v in enumerate(row_values):
        header = str(headers[i]).strip() if i < len(headers) and headers[i] else f"col_{i}"
        raw_data[header] = v

    record = ParsedInvoiceRecord()
    record.source = _determine_source(category)
    record.source_sheet = sheet_name
    record.sheet_category = category
    record.source_row = row_number
    record.raw_data = raw_data

    # Supplier name
    supplier_raw = _get("supplier_name")
    record.supplier_name = str(supplier_raw).strip() if supplier_raw else ""

    # GSTIN
    gstin_raw = _get("gstin")
    orig, norm, valid = normalize_gstin(gstin_raw)
    record.gstin_original = orig
    record.gstin_normalized = norm
    record.gstin_valid = valid

    # Invoice number.  B2BA exposes both original and revised identity; the
    # revised identity is the one that participates in reconciliation.
    inv_raw = _get("revised_invoice_number") if category == SheetCategory.B2BA else _get("invoice_number")
    record.invoice_number = normalize_invoice_number(inv_raw)

    # Invoice date
    date_raw = _get("revised_invoice_date") if category == SheetCategory.B2BA else _get("invoice_date")
    parsed_date, date_orig = parse_date(date_raw)
    record.invoice_date = parsed_date
    record.invoice_date_original = date_orig

    # Voucher type
    vtype_raw = _get("voucher_type")
    if category in (SheetCategory.CDNR, SheetCategory.CDNRA) and "note_type" in col_map:
        vtype_raw = _get("note_type")
    record.voucher_type, record.voucher_type_original = _determine_voucher_type(
        vtype_raw, category
    )

    # Monetary values
    taxable_raw = _get("taxable_value")
    record.taxable_value_present = not (
        taxable_raw is None
        or (isinstance(taxable_raw, str) and not taxable_raw.strip())
    )
    record.taxable_value = _money("taxable_value")
    record.igst = _money("igst")
    record.cgst = _money("cgst")
    record.sgst = _money("sgst")
    cess_raw = _get("cess")
    record.cess_present = not (
        cess_raw is None
        or (isinstance(cess_raw, str) and not cess_raw.strip())
    )
    record.cess = _money("cess", cess_raw)
    total_invoice_raw = _get("total_invoice_value")
    if (
        category in (SheetCategory.CDNR, SheetCategory.CDNRA)
        and (
            total_invoice_raw is None
            or (isinstance(total_invoice_raw, str) and not total_invoice_raw.strip())
        )
    ):
        total_invoice_raw = _get("note_value")
    record.total_invoice_value_present = not (
        total_invoice_raw is None
        or (isinstance(total_invoice_raw, str) and not total_invoice_raw.strip())
    )
    record.total_invoice_value = _money("total_invoice_value", total_invoice_raw)
    record.additional_cost = _money("additional_cost")
    gross_raw = _get("gross_total")
    # Backward compatibility for persisted sessions created before Gross Total
    # became its own Accounts mapping.
    if (gross_raw is None or (isinstance(gross_raw, str) and not gross_raw.strip())) and record.source == RecordSource.TALLY:
        gross_raw = total_invoice_raw
    record.gross_total_present = not (
        gross_raw is None or (isinstance(gross_raw, str) and not gross_raw.strip())
    )
    record.gross_total = _money("gross_total", gross_raw)
    record.round_off = _money("round_off")

    # This Tally DayBook export leaves Value blank when the taxable amount is
    # carried by a named purchase ledger. In that exact structural format,
    # derive the missing taxable value from the mutually exclusive ledgers.
    daybook_signature = {
        "Column B (blank header)", "Column C (blank header)",
        "Supplier Invoice No.", "GSTIN/UIN", "Value", "Gross Total",
        "CGST", "SGST", "I-GST",
    }
    if daybook_signature.issubset(set(str(header) for header in headers)) and not record.taxable_value_present:
        taxable_ledgers = {
            "Purchase 18%", "Pur 5%", "Number Palat  18%", "Tractors Pur 5%",
            "Frieght 18% GST Paid", "Air Conditioner", "SERVICE INCOME 18%",
        }
        ledger_values = []
        for header in taxable_ledgers:
            raw_ledger = raw_data.get(header)
            try:
                ledger_values.append(parse_money(raw_ledger))
            except (ValueError, ArithmeticError):
                record.financial_parse_issues.append(FinancialParseIssue(
                    field_name="taxable_value",
                    raw_value=str(raw_ledger),
                    reason=f"Taxable ledger '{header}' contains an invalid monetary value: {raw_ledger!r}",
                ))
        record.taxable_value = sum(ledger_values, Decimal("0.00"))

    # GST-specific
    rate_raw = _get("gst_rate")
    try:
        record.gst_rate = parse_rate(rate_raw)
    except (ValueError, ArithmeticError):
        record.financial_parse_issues.append(FinancialParseIssue(
            field_name="gst_rate", raw_value=str(rate_raw),
            reason=f"gst_rate contains an invalid percentage/rate: {rate_raw!r}",
        ))
    record.place_of_supply = str(_get("place_of_supply") or "").strip()
    record.reverse_charge = str(_get("reverse_charge") or "").strip()

    # CDNR-specific
    record.note_type = str(_get("note_type") or "").strip()
    note_num_raw = _get("revised_note_number") if category == SheetCategory.CDNRA else _get("note_number")
    if note_num_raw:
        note_info = normalize_invoice_number(note_num_raw)
        record.note_number = note_info.original
        # Also populate invoice_number with the note for unified matching
        if not record.invoice_number.original:
            record.invoice_number = note_info

    note_date_raw = _get("revised_note_date") if category == SheetCategory.CDNRA else _get("note_date")
    if note_date_raw:
        nd, _ = parse_date(note_date_raw)
        record.note_date = nd

    orig_inv_raw = _get("original_invoice_number")
    record.original_invoice_number = str(orig_inv_raw).strip() if orig_inv_raw else ""

    amendment_original_raw = _get("amendment_original_number")
    record.amendment_original_number = str(amendment_original_raw).strip() if amendment_original_raw else ""
    amendment_original_date_raw = _get("amendment_original_date")
    if amendment_original_date_raw is not None:
        record.amendment_original_date, record.amendment_original_date_original = parse_date(
            amendment_original_date_raw
        )
    record.amendment_period = str(_get("amendment_period") or "").strip()
    amendment_date_raw = _get("amendment_date")
    if amendment_date_raw is not None:
        record.amendment_date, _ = parse_date(amendment_date_raw)

    # Skip rows that have no meaningful data (no supplier, no GSTIN, no invoice)
    has_identity = (
        record.supplier_name
        or record.gstin_normalized
        or record.invoice_number.original
    )
    if daybook_signature.issubset(set(str(header) for header in headers)):
        has_identity = bool(record.gstin_normalized or record.invoice_number.original)
    if not has_identity:
        return None

    return record


# ---------------------------------------------------------------------------
# Excel file parsing
# ---------------------------------------------------------------------------

def _read_sheet_rows(ws: Any) -> List[List[Any]]:
    """Read all rows from an openpyxl worksheet as lists of values."""
    rows = []
    for row in ws.iter_rows():
        rows.append([cell.value for cell in row])
    return rows


def _read_xls_sheet_rows(sheet: Any, datemode: int) -> List[List[Any]]:
    """Read an xlrd sheet using values compatible with the openpyxl path."""
    rows: List[List[Any]] = []
    for row_index in range(sheet.nrows):
        values: List[Any] = []
        for cell in sheet.row(row_index):
            if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                value = None
            elif cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate.xldate_as_datetime(cell.value, datemode)
                except (ValueError, OverflowError):
                    value = cell.value
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                value = bool(cell.value)
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                value = None
            elif cell.ctype == xlrd.XL_CELL_NUMBER and cell.value.is_integer():
                # openpyxl returns integral numeric cells as ints. Preserve that
                # behavior so numeric invoice IDs never gain a trailing ".0".
                value = int(cell.value)
            else:
                value = cell.value
            values.append(value)
        rows.append(values)
    return rows


def parse_excel_file(
    file_path: str,
    source_hint: Optional[RecordSource] = None,
    sheet_name_override: Optional[str] = None,
    mapped_headers: Optional[Dict[str, str]] = None,
) -> FileInspectionResult:
    """
    Parse an Excel .xls or .xlsx file into the same canonical row structure.
    
    Args:
        file_path: Path to the Excel file.
        source_hint: Optional hint about whether this is a Tally or GST file.
        sheet_name_override: Optional specific sheet to parse.
        mapped_headers: Optional mapping overrides.
    
    Returns:
        FileInspectionResult with detected sheets and parsed records.
    """
    result = FileInspectionResult()
    result.file_name = os.path.basename(file_path)
    result.file_type = Path(file_path).suffix.lstrip(".").lower()
    result.file_size_bytes = os.path.getsize(file_path)

    is_legacy_xls = Path(file_path).suffix.lower() == ".xls"
    try:
        if is_legacy_xls:
            wb = xlrd.open_workbook(file_path, on_demand=True)
            sheet_names = wb.sheet_names()
        else:
            # Keep the established .xlsx reader and its behavior unchanged.
            wb = openpyxl.load_workbook(file_path, data_only=True, read_only=False)
            sheet_names = wb.sheetnames
    except Exception as e:
        result.errors.append(f"Failed to open Excel file: {e}")
        return result

    for sheet_name in sheet_names:
        if sheet_name_override and sheet_name != sheet_name_override:
            continue
        if is_legacy_xls:
            rows = _read_xls_sheet_rows(wb.sheet_by_name(sheet_name), wb.datemode)
        else:
            rows = _read_sheet_rows(wb[sheet_name])

        # Detect sheet structure
        sheet_result = detect_sheet(sheet_name, rows)
        header_values = sheet_result.headers
        data_start = max(0, sheet_result.data_start_row - 1)
        result.persisted_sheets.append(PersistedSheetRows(
            sheet_name=sheet_name,
            category=sheet_result.category,
            headers=header_values,
            rows=rows[data_start:],
            row_numbers=list(range(data_start + 1, len(rows) + 1)),
        ))
        
        if mapped_headers is not None:
            # Override or add column mappings
            existing_semantics = {m.semantic_name: m for m in sheet_result.column_mappings}
            
            for semantic_name, header_val in mapped_headers.items():
                if header_val is None or header_val == "" or header_val == "-- Not mapped --":
                    if semantic_name in existing_semantics:
                        sheet_result.column_mappings.remove(existing_semantics[semantic_name])
                        del existing_semantics[semantic_name]
                else:
                    if semantic_name in existing_semantics:
                        existing_semantics[semantic_name].detected_header = header_val
                    else:
                        sheet_result.column_mappings.append(
                            ColumnMapping(semantic_name=semantic_name, detected_header=header_val, confidence=1.0)
                        )

        # If sheet has no data and no recognized headers, skip
        if sheet_result.total_data_rows == 0 and sheet_result.confidence < 0.1 and not mapped_headers:
            result.warnings.append(
                f"Sheet '{sheet_name}' has no data and unrecognized headers. Skipped."
            )
            # Still add sheet metadata for inspection
            result.sheets.append(sheet_result)
            continue

        result.sheets.append(sheet_result)

        # Build column index map
        col_map = build_column_index_map(
            header_values,
            sheet_result.column_mappings,
        )
        _apply_amendment_column_map(header_values, col_map, sheet_result.category)

        # Parse data rows
        data_start = max(0, sheet_result.data_start_row - 1)

        for row_idx in range(data_start, len(rows)):
            row_values = rows[row_idx]
            record = _row_to_record(
                row_values=row_values,
                headers=header_values,
                col_map=col_map,
                category=sheet_result.category,
                sheet_name=sheet_name,
                row_number=row_idx + 1,  # 1-indexed for user display
            )
            if record is not None:
                # Apply source hint override if provided
                if source_hint is not None:
                    record.source = source_hint
                result.records.append(record)

    if is_legacy_xls:
        wb.release_resources()
    else:
        wb.close()
    return result


def reparse_persisted_file(
    inspected: FileInspectionResult,
    source_hint: Optional[RecordSource] = None,
    sheet_name_override: Optional[str] = None,
    mapped_headers: Optional[Dict[str, Optional[str]]] = None,
) -> FileInspectionResult:
    """Rebuild parsed records from the canonical rows persisted at inspect time."""
    if not inspected.persisted_sheets:
        if mapped_headers:
            raise ValueError("This session predates durable mapping data. Please inspect the files again.")
        return copy.deepcopy(inspected)
    result = FileInspectionResult(
        file_name=inspected.file_name,
        file_type=inspected.file_type,
        file_size_bytes=inspected.file_size_bytes,
        errors=list(inspected.errors),
        warnings=list(inspected.warnings),
        persisted_sheets=inspected.persisted_sheets,
    )
    sheets_by_name = {sheet.sheet_name: sheet for sheet in inspected.sheets}

    for source_sheet in inspected.persisted_sheets:
        if sheet_name_override and source_sheet.sheet_name != sheet_name_override:
            continue
        detected = sheets_by_name.get(source_sheet.sheet_name)
        if detected is None:
            continue
        sheet_result = copy.deepcopy(detected)
        existing = {mapping.semantic_name: mapping for mapping in sheet_result.column_mappings}
        if mapped_headers is not None:
            for semantic_name, header_value in mapped_headers.items():
                if header_value is None or header_value == "" or header_value == "-- Not mapped --":
                    mapping = existing.pop(semantic_name, None)
                    if mapping is not None:
                        sheet_result.column_mappings.remove(mapping)
                elif semantic_name in existing:
                    existing[semantic_name].detected_header = header_value
                else:
                    mapping = ColumnMapping(
                        semantic_name=semantic_name,
                        detected_header=header_value,
                        confidence=1.0,
                    )
                    sheet_result.column_mappings.append(mapping)
                    existing[semantic_name] = mapping

        result.sheets.append(sheet_result)
        column_map = build_column_index_map(source_sheet.headers, sheet_result.column_mappings)
        _apply_amendment_column_map(source_sheet.headers, column_map, source_sheet.category)
        for row_values, row_number in zip(source_sheet.rows, source_sheet.row_numbers):
            record = _row_to_record(
                row_values=row_values,
                headers=source_sheet.headers,
                col_map=column_map,
                category=source_sheet.category,
                sheet_name=source_sheet.sheet_name,
                row_number=row_number,
            )
            if record is not None:
                if source_hint is not None and inspected.file_type in ("xlsx", "xls"):
                    record.source = source_hint
                result.records.append(record)
    return result


# ---------------------------------------------------------------------------
# JSON file parsing
# ---------------------------------------------------------------------------

def _json_records_to_parsed(
    records: List[Dict[str, Any]],
    sheet_name: str,
    category: SheetCategory,
) -> Tuple[List[ParsedInvoiceRecord], SheetDetectionResult]:
    """
    Convert a list of JSON record dicts into ParsedInvoiceRecords.
    
    Uses the same header/column detection logic — treats JSON keys as headers.
    """
    if not records:
        return ([], SheetDetectionResult(sheet_name=sheet_name, category=category))

    # Treat JSON keys as the header row
    all_keys: List[str] = []
    seen = set()
    for rec in records:
        for k in rec.keys():
            if k not in seen:
                all_keys.append(k)
                seen.add(k)

    # Detect columns from JSON keys
    mappings = detect_columns(all_keys)
    col_map: Dict[str, int] = {}
    for m in mappings:
        for idx, k in enumerate(all_keys):
            if str(k).strip() == m.detected_header:
                col_map[m.semantic_name] = idx
                break
    _apply_amendment_column_map(all_keys, col_map, category)

    # Build sheet detection result
    sheet_result = SheetDetectionResult(
        sheet_name=sheet_name,
        category=category,
        header_row=0,
        data_start_row=1,
        total_data_rows=len(records),
        headers=all_keys,
        column_mappings=mappings,
        confidence=len(mappings) / max(1, len(all_keys)),
    )

    # Parse each record
    parsed: List[ParsedInvoiceRecord] = []
    for row_num, rec in enumerate(records, 1):
        row_values = [rec.get(k) for k in all_keys]
        record = _row_to_record(
            row_values=row_values,
            headers=all_keys,
            col_map=col_map,
            category=category,
            sheet_name=sheet_name,
            row_number=row_num,
        )
        if record is not None:
            parsed.append(record)

    return (parsed, sheet_result)


def _guess_json_category(sheet_name: str) -> SheetCategory:
    """Guess the sheet category from a JSON object key name."""
    name_lower = sheet_name.strip().lower()
    mapping = {
        "purchase register": SheetCategory.PURCHASE_REGISTER,
        "purchase": SheetCategory.PURCHASE_REGISTER,
        "purchases": SheetCategory.PURCHASE_REGISTER,
        "debit note register": SheetCategory.DEBIT_NOTE_REGISTER,
        "debit note": SheetCategory.DEBIT_NOTE_REGISTER,
        "debit notes": SheetCategory.DEBIT_NOTE_REGISTER,
        "b2b": SheetCategory.B2B,
        "b2ba": SheetCategory.B2BA,
        "cdnr": SheetCategory.CDNR,
        "notes": SheetCategory.CDNR,
        "note": SheetCategory.CDNR,
        "cdnra": SheetCategory.CDNRA,
        "party": SheetCategory.PARTY,
    }
    return mapping.get(name_lower, SheetCategory.UNKNOWN)


def parse_json_file(file_path: str) -> FileInspectionResult:
    """
    Parse a JSON file.
    
    Supports two formats:
    1. Direct array: [{...}, {...}, ...]
    2. Workbook-like object: {"Purchase Register": [...], "b2b": [...], ...}
    
    Returns:
        FileInspectionResult with detected structure and parsed records.
    """
    result = FileInspectionResult()
    result.file_name = os.path.basename(file_path)
    result.file_type = "json"
    result.file_size_bytes = os.path.getsize(file_path)

    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, UnicodeDecodeError, RecursionError) as e:
        result.errors.append(f"Failed to parse JSON: {e}")
        return result

    if isinstance(data, list):
        if not all(isinstance(record, dict) for record in data):
            result.errors.append("JSON arrays must contain objects.")
            return result
        # Direct array format
        records, sheet = _json_records_to_parsed(
            data, "data", SheetCategory.UNKNOWN
        )
        result.sheets.append(sheet)
        result.records.extend(records)
        result.persisted_sheets.append(PersistedSheetRows(
            sheet_name="data", category=SheetCategory.UNKNOWN,
            headers=sheet.headers,
            rows=[[record.get(key) for key in sheet.headers] for record in data],
            row_numbers=list(range(1, len(data) + 1)),
        ))

    elif isinstance(data, dict):
        # Workbook-like object
        found_array = False
        for key, value in data.items():
            if isinstance(value, list):
                found_array = True
                if not all(isinstance(record, dict) for record in value):
                    result.errors.append(f"JSON sheet '{key}' must contain objects.")
                    continue
                category = _guess_json_category(key)
                records, sheet = _json_records_to_parsed(value, key, category)
                result.sheets.append(sheet)
                result.records.extend(records)
                result.persisted_sheets.append(PersistedSheetRows(
                    sheet_name=key, category=category,
                    headers=sheet.headers,
                    rows=[[record.get(header) for header in sheet.headers] for record in value],
                    row_numbers=list(range(1, len(value) + 1)),
                ))
            else:
                result.warnings.append(
                    f"JSON key '{key}' is not an array. Skipped."
                )
        if not found_array:
            result.errors.append("JSON object must contain at least one array of records.")
    else:
        result.errors.append("JSON root must be an array or object.")

    return result


# ---------------------------------------------------------------------------
# Unified entry point
# ---------------------------------------------------------------------------

def parse_file(
    file_path: str,
    source_hint: Optional[RecordSource] = None,
) -> FileInspectionResult:
    """
    Parse any supported file (.xls, .xlsx or .json).
    
    Args:
        file_path: Path to the file.
        source_hint: Optional hint about whether this is a Tally or GST file.
    
    Returns:
        FileInspectionResult.
    
    Raises:
        ValueError: If file format is not supported.
    """
    ext = Path(file_path).suffix.lower()

    if ext in (".xls", ".xlsx"):
        return parse_excel_file(file_path, source_hint)
    elif ext == ".json":
        return parse_json_file(file_path)
    else:
        raise ValueError(
            f"Unsupported file format: {ext}. Supported formats: .xls, .xlsx and .json"
        )
