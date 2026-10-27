import json
import os
import pytest
import tempfile
from decimal import Decimal
from datetime import date
from openpyxl import load_workbook

from backend.app.models.records import ParsedInvoiceRecord, RecordSource, SheetCategory, VoucherType
from backend.app.models.reconciliation import (
    ReconciliationSummary, ReconciliationResult, ReconciliationStatus, MatchMethod, AuditTrail, TaxComparison,
    AggregatedGSTInvoice
)
from backend.app.services.notes_reconciliation_engine import (
    NotesReconciliationSummary, AggregatedGSTCreditNote
)
from backend.app.services.json_exporter import generate_json_report
from backend.app.services.excel_exporter import (
    export_notes_to_excel, export_purchase_to_excel, export_to_excel,
)
from backend.app.services.summary_generator import generate_text_summary

D = Decimal


@pytest.fixture(scope="module")
def mock_summary():
    s = ReconciliationSummary()
    s.total_tally_invoices = 10
    s.total_gst_invoices = 9
    
    r1 = ReconciliationResult(
        supplier_gstin="06DEMOX1234A1Z5", supplier_tally="Supplier 1", supplier_gst="Supplier 1",
        invoice_number_tally="INV001", invoice_number_gst="INV001", invoice_date=date(2025, 4, 1),
        taxable_value=D("1000.00"), cgst_diff=D("0.00"), sgst_diff=D("0.00"), igst_diff=D("0.00"),
        status=ReconciliationStatus.PERFECTLY_MATCHED, match_method=MatchMethod.EXACT_INVOICE,
        remarks="Exact Match"
    )
    r1.audit = AuditTrail()
    
    r2 = ReconciliationResult(
        supplier_gstin="06DEMOX1234A1Z5", supplier_tally="Supplier 2",
        invoice_number_tally="INV002", invoice_date=date(2025, 4, 2),
        taxable_value=D("2000.00"),
        status=ReconciliationStatus.MISSING_IN_GSTR, match_method=MatchMethod.NO_MATCH,
        remarks="Missing in GSTR-2B"
    )
    r2.audit = AuditTrail(tally_rows=[ParsedInvoiceRecord(source=RecordSource.TALLY, cgst=D("180"), sgst=D("180"), igst=D("0"))])

    r3 = ReconciliationResult(
        supplier_gstin="06DEMOX1234A1Z5", supplier_tally="Supplier 3", supplier_gst="Supplier 3",
        invoice_number_tally="INV003", invoice_number_gst="INV003", invoice_date=date(2025, 4, 3),
        taxable_value=D("3000.00"), cgst_diff=D("2.00"), sgst_diff=D("-2.00"), igst_diff=D("0.00"),
        status=ReconciliationStatus.MAJOR_DIFFERENCE, match_method=MatchMethod.EXACT_INVOICE,
        remarks="Tax Diff"
    )
    r3.audit = AuditTrail(
        cgst_comparison=TaxComparison("CGST", D("272.00"), D("270.00"), D("2.00"), abs_difference=D("2.00")),
        sgst_comparison=TaxComparison("SGST", D("268.00"), D("270.00"), D("-2.00"), abs_difference=D("2.00")),
        igst_comparison=TaxComparison("IGST", D("0"), D("0"), D("0"), abs_difference=D("0"))
    )

    r4 = ReconciliationResult(
        supplier_gstin="06DEMOX1234A1Z5", supplier_tally="Supplier 4", supplier_gst="Supplier 4",
        invoice_number_tally="TALLY4587", invoice_number_gst="GST/4587", invoice_date=date(2025, 4, 4),
        taxable_value=D("4000.00"), cgst_diff=D("0.00"), sgst_diff=D("0.00"), igst_diff=D("0.00"),
        status=ReconciliationStatus.PERFECTLY_MATCHED, match_method=MatchMethod.INVOICE_LAST_4,
        remarks="Matched using last 4"
    )
    r4.audit = AuditTrail()
    
    r5 = ReconciliationResult(
        supplier_gstin="06DEMOX1234A1Z5", supplier_tally="Supplier Ambig", supplier_gst="Supplier Ambig",
        invoice_number_tally="AMBIG001", invoice_number_gst="", invoice_date=date(2025, 4, 5),
        taxable_value=D("5000.00"), cgst_diff=D("0.00"), sgst_diff=D("0.00"), igst_diff=D("0.00"),
        status=ReconciliationStatus.NEEDS_REVIEW, match_method=MatchMethod.AMBIGUOUS,
        remarks="Ambiguous"
    )
    r5.audit = AuditTrail()

    s.results = [r1, r2, r3, r4, r5]
    
    s.not_in_tally_records = [
        AggregatedGSTInvoice(
            gstin="06FFFFK9103A1Z0", supplier_name="Supplier NIT", invoice_number_original="NIT001",
            invoice_date=date(2025, 4, 6), taxable_value=D("6000.00")
        )
    ]
    
    s.compute_counts()
    return s


@pytest.fixture(scope="module")
def mock_notes():
    n = NotesReconciliationSummary()
    n.total_tally_debit_notes = 1
    n.total_gst_credit_notes = 1
    
    r1 = ReconciliationResult(
        supplier_gstin="06DEMOX1234A1Z5", supplier_tally="Supplier 5", supplier_gst="Supplier 5",
        invoice_number_tally="DN001", invoice_number_gst="DN001", invoice_date=date(2025, 4, 5),
        taxable_value=D("500.00"), cgst_diff=D("0.00"), sgst_diff=D("0.00"), igst_diff=D("0.00"),
        status=ReconciliationStatus.PERFECTLY_MATCHED, match_method=MatchMethod.EXACT_INVOICE,
        remarks="Exact Match"
    )
    r1.audit = AuditTrail(tally_rows=[ParsedInvoiceRecord(source=RecordSource.TALLY, invoice_date=date(2025,4,5))], tally_taxable=D("500.00"))
    n.results = [r1]
    
    agg = AggregatedGSTCreditNote(
        gstin="06DEMOX1234A1Z5", supplier_name="Supplier 6", note_number_original="CN002",
        note_date=date(2025, 4, 6), taxable_value=D("100.00")
    )
    n.not_in_tally_notes = [agg]
    
    n.compute_counts()
    return n


@pytest.fixture(scope="module")
def generated_excel_path(mock_summary, mock_notes):
    with tempfile.TemporaryDirectory() as tmpdir:
        file_path = os.path.join(tmpdir, "test_report.xlsx")
        export_to_excel(file_path, mock_summary, mock_notes, "April 2025")
        yield file_path


@pytest.fixture(scope="module")
def parsed_json(mock_summary, mock_notes):
    json_str = generate_json_report(mock_summary, mock_notes, "April 2025")
    return json.loads(json_str)


@pytest.fixture
def invoice_mismatch_summary():
    tally = ParsedInvoiceRecord(
        source=RecordSource.TALLY, supplier_name="Mismatch Supplier",
        gstin_original="06SYNTH1234A1Z5", gstin_normalized="06SYNTH1234A1Z5",
        invoice_date=date(2025, 5, 1), taxable_value=D("1000.00"),
        igst=D("0.00"), cgst=D("90.00"), sgst=D("90.00"),
    )
    gst = ParsedInvoiceRecord(
        source=RecordSource.GST, supplier_name="Mismatch Supplier",
        gstin_original="06SYNTH1234A1Z5", gstin_normalized="06SYNTH1234A1Z5",
        invoice_date=date(2025, 5, 2), taxable_value=D("1000.00"),
        igst=D("0.00"), cgst=D("90.00"), sgst=D("90.00"),
    )
    result = ReconciliationResult(
        supplier_gstin="06SYNTH1234A1Z5", supplier_tally="Mismatch Supplier",
        supplier_gst="Mismatch Supplier", invoice_number_tally="SGST-5725",
        invoice_number_gst="SGST-5728", invoice_date=date(2025, 5, 1),
        tally_date=date(2025, 5, 1), gstr_date=date(2025, 5, 2),
        taxable_value=D("1000.00"), status=ReconciliationStatus.INVOICE_NUMBER_MISMATCH,
        match_method=MatchMethod.INVOICE_MISMATCH, invoice_mismatch=True,
        date_mismatch=True, date_difference_days=1, remarks="Invoice numbers differ",
        audit=AuditTrail(
            tally_rows=[tally], gst_rows=[gst], tally_taxable=D("1000.00"),
            gst_taxable=D("1000.00"), explanation="Exact GSTIN, date and financial evidence",
            igst_comparison=TaxComparison("IGST", D("0.00"), D("0.00"), D("0.00")),
            cgst_comparison=TaxComparison("CGST", D("90.00"), D("90.00"), D("0.00")),
            sgst_comparison=TaxComparison("SGST", D("90.00"), D("90.00"), D("0.00")),
        ),
    )
    result.api_result_id = "result-mismatch-1"
    summary = ReconciliationSummary(total_tally_invoices=1, total_gst_invoices=1, results=[result])
    summary.compute_counts()
    return summary


def _notes_with_compatibility_duplicate(*, consumed=False):
    agg = AggregatedGSTCreditNote(
        gstin="06BBBBB0000B1Z5", supplier_name="Unmatched Note Supplier",
        note_number_original="CN-DUP-1", note_number_normalized="CNDUP1",
        note_date=date(2025, 6, 1), taxable_value=D("250.00"), consumed=consumed,
    )
    synthetic = ReconciliationResult(
        supplier_gstin=agg.gstin, supplier_gst=agg.supplier_name,
        invoice_number_gst=agg.note_number_original, invoice_date=agg.note_date,
        taxable_value=agg.taxable_value, status=ReconciliationStatus.NOT_IN_TALLY,
        match_method=MatchMethod.NO_MATCH,
        remarks="GST credit note not matched to any Tally debit note",
        audit=AuditTrail(
            gst_invoice_original=agg.note_number_original,
            gst_invoice_normalized=agg.note_number_normalized,
            gst_gstin=agg.gstin, gst_supplier=agg.supplier_name,
            gst_taxable=agg.taxable_value,
        ),
    )
    notes = NotesReconciliationSummary(
        total_gst_credit_notes=1,
        results=[] if consumed else [synthetic],
        not_in_tally_notes=[agg],
    )
    notes.compute_counts()
    return notes


class TestExcelExporter:
    def test_purchase_export_is_categorized_and_priority_ordered(self, tmp_path, mock_summary):
        path = tmp_path / "categorized-purchase.xlsx"
        export_purchase_to_excel(str(path), mock_summary, "April 2025")
        wb = load_workbook(path)
        assert wb.sheetnames[:13] == [
            "Summary", "All Results", "Perfect", "Minor", "Major", "Invoice Mismatch",
            "Date Mismatch", "Missing in GSTR", "Missing in Accounts", "Needs Review",
            "Taxable Missing", "Potential Duplicate", "Cancelled Vouchers",
        ]
        assert wb["Summary"]["A1"].value == "Purchase Reconciliation Report"
        assert wb["All Results"]["A1"].value == "#"
        assert wb["All Results"]["AB2"].value == "Needs Review / Ambiguous"
        assert wb["All Results"]["AI1"].value == "Accounts Source Rows"
        assert wb["Perfect"]["AB2"].value == "Perfectly Matched"
        major_headers = [cell.value for cell in wb["Major"][1]]
        assert wb["Major"].cell(2, major_headers.index("Status") + 1).value == "Major Difference"
        assert wb["Reconciliation Summary"].sheet_state == "hidden"
        assert wb["Invoice No Remarks"].sheet_state == "hidden"

    def test_missing_purchase_sheets_include_only_the_available_source_columns(self, tmp_path):
        accounts_row = ParsedInvoiceRecord(
            source=RecordSource.TALLY, source_sheet="Purchase Register", source_row=12,
            supplier_name="Accounts Supplier", gstin_normalized="06AAAAA0000A1Z5",
            invoice_date=date(2025, 4, 1), taxable_value=D("1000.00"),
            gross_total=D("1180.00"), igst=D("180.00"),
        )
        missing_in_gstr = ReconciliationResult(
            supplier_gstin="06AAAAA0000A1Z5", supplier_tally="Accounts Supplier",
            invoice_number_tally="ACC-1", status=ReconciliationStatus.MISSING_IN_GSTR,
            match_method=MatchMethod.NO_MATCH, remarks="Missing in GSTR-2B",
            audit=AuditTrail(
                tally_rows=[accounts_row], tally_invoice_normalized="ACC1",
                tally_gross_total=D("1180.00"), explanation="No matching GSTR invoice",
            ),
        )
        gstr_row = ParsedInvoiceRecord(
            source=RecordSource.GST, source_sheet="invoice", source_row=34,
            sheet_category=SheetCategory.B2B, supplier_name="GSTR Supplier",
            gstin_normalized="07BBBBB0000B1Z3", invoice_date=date(2025, 4, 2),
            taxable_value=D("2000.00"), total_invoice_value=D("2360.00"),
            igst=D("360.00"),
        )
        missing_in_accounts = ReconciliationResult(
            supplier_gstin="07BBBBB0000B1Z3", supplier_gst="GSTR Supplier",
            invoice_number_gst="GST-2", status=ReconciliationStatus.NOT_IN_TALLY,
            match_method=MatchMethod.NO_MATCH, remarks="Not in Accounts",
            audit=AuditTrail(
                gst_rows=[gstr_row], gst_invoice_normalized="GST2",
                gst_invoice_value=D("2360.00"), explanation="No matching Accounts invoice",
            ),
        )
        summary = ReconciliationSummary(
            total_tally_invoices=1, total_gst_invoices=1,
            results=[missing_in_gstr, missing_in_accounts],
        )
        summary.compute_counts()
        path = tmp_path / "source-specific-missing-sheets.xlsx"

        export_purchase_to_excel(str(path), summary, "April 2025")

        wb = load_workbook(path)
        accounts_sheet = wb["Missing in GSTR"]
        accounts_headers = [cell.value for cell in accounts_sheet[1]]
        assert accounts_headers == [
            "#", "Supplier GSTIN (Accounts)", "Supplier (Accounts)",
            "Invoice No. (Accounts)", "Accounts Date",
            "Taxable Value (Accounts)", "Gross Total (Accounts)",
            "IGST (Accounts)", "CGST (Accounts)", "SGST (Accounts)",
            "Match Method", "Status", "Remarks",
        ]
        assert accounts_sheet.cell(2, 4).value == "ACC-1"
        assert "No matching GSTR invoice" in accounts_sheet.cell(2, 13).value
        assert accounts_sheet.auto_filter.ref == "A1:M2"

        gstr_sheet = wb["Missing in Accounts"]
        gstr_headers = [cell.value for cell in gstr_sheet[1]]
        assert gstr_headers == [
            "#", "Supplier GSTIN (GSTR-2B)", "Supplier (GSTR-2B)",
            "Invoice No. (GSTR-2B)", "GSTR Date",
            "Taxable Value (GSTR-2B)", "Invoice Value (GSTR-2B)",
            "IGST (GSTR-2B)", "CGST (GSTR-2B)", "SGST (GSTR-2B)",
            "Match Method", "Status", "Remarks", "GSTR Category",
        ]
        assert gstr_sheet.cell(2, 4).value == "GST-2"
        assert "No matching Accounts invoice" in gstr_sheet.cell(2, 13).value
        assert gstr_sheet.cell(2, 14).value == "b2b"
        assert gstr_sheet.auto_filter.ref == "A1:N2"

    def test_purchase_review_sheets_use_compact_columns_and_plain_counts(self, tmp_path, invoice_mismatch_summary):
        path = tmp_path / "compact-purchase.xlsx"
        export_purchase_to_excel(str(path), invoice_mismatch_summary, "April 2025")

        wb = load_workbook(path)
        removed_headers = {
            "Supplier (Accounts)", "Invoice No. (Accounts)",
            "Normalized Invoice (GSTR-2B)", "GSTR Date", "Evidence",
            "GSTR Source Sheet", "GSTR Source Rows",
            "Accounts Source Sheet", "Accounts Source Rows",
        }
        for sheet_name in ("Minor", "Major", "Invoice Mismatch"):
            headers = {cell.value for cell in wb[sheet_name][1]}
            assert headers.isdisjoint(removed_headers)

        mismatch_sheet = wb["Invoice Mismatch"]
        mismatch_headers = [cell.value for cell in mismatch_sheet[1]]
        remarks = mismatch_sheet.cell(2, mismatch_headers.index("Remarks") + 1).value
        assert "Invoice numbers differ" in remarks
        assert "Exact GSTIN, date and financial evidence" in remarks

        summary_sheet = wb["Summary"]
        assert "₹" not in summary_sheet["B7"].number_format
        assert summary_sheet["B7"].number_format == "#,##0"
        assert wb["All Results"]["A2"].number_format == "#,##0"
        assert wb["All Results"].column_dimensions["A"].width >= 8

    def test_potential_duplicate_prunes_empty_side_and_source_columns(self, tmp_path):
        accounts_row = ParsedInvoiceRecord(
            source=RecordSource.TALLY, source_sheet="Purchase Register", source_row=522,
            supplier_name="Duplicate Supplier", gstin_normalized="06AAAAA0000A1Z5",
            invoice_date=date(2025, 4, 1), taxable_value=D("1000.00"),
            gross_total=D("1180.00"), cgst=D("90.00"), sgst=D("90.00"),
        )
        duplicate = ReconciliationResult(
            supplier_gstin="06AAAAA0000A1Z5", supplier_tally="Duplicate Supplier",
            invoice_number_tally="DUP-1", tally_date=date(2025, 4, 1),
            status=ReconciliationStatus.POTENTIAL_DUPLICATE,
            match_method=MatchMethod.NO_MATCH, remarks="Potential duplicate",
            audit=AuditTrail(
                tally_rows=[accounts_row], tally_invoice_normalized="DUP1",
                tally_taxable=D("1000.00"), tally_gross_total=D("1180.00"),
                explanation="Two Accounts records share GSTIN and invoice",
            ),
        )
        summary = ReconciliationSummary(total_tally_invoices=2, results=[duplicate])
        summary.compute_counts()
        path = tmp_path / "potential-duplicate.xlsx"

        export_purchase_to_excel(str(path), summary, "April 2025")

        wb = load_workbook(path)
        sheet = wb["Potential Duplicate"]
        headers = [cell.value for cell in sheet[1]]
        assert "Supplier (Accounts)" in headers
        assert "Supplier (GSTR-2B)" not in headers
        assert "GSTR Source Sheet" not in headers
        assert "GSTR Source Rows" not in headers
        assert "Accounts Source Sheet" not in headers
        assert "Accounts Source Rows" not in headers
        assert "Evidence" not in headers
        for col_idx in range(1, sheet.max_column + 1):
            assert sheet.cell(2, col_idx).value not in (None, "", "—")

    def test_purchase_and_notes_workbooks_are_separate(self, tmp_path, mock_summary, mock_notes):
        purchase_path = tmp_path / "purchase.xlsx"
        notes_path = tmp_path / "notes.xlsx"
        export_purchase_to_excel(str(purchase_path), mock_summary, "April 2025")
        export_notes_to_excel(str(notes_path), mock_notes, "April 2025")

        purchase = load_workbook(purchase_path)
        notes = load_workbook(notes_path)
        assert "Notes" not in purchase.sheetnames
        assert notes.sheetnames == [
            "Summary", "Perfect", "Minor", "Tax Mismatch", "Missing in GSTR",
            "Missing in Debit Notes", "Needs Review",
        ]
        assert notes["Summary"]["B2"].value == "April 2025"
        assert notes["Perfect"].cell(2, 5).value == "DN001"
        assert "Reconciliation Summary" not in notes.sheetnames

    def test_workbook_generation_and_reopen(self, generated_excel_path):
        assert os.path.exists(generated_excel_path)
        wb = load_workbook(generated_excel_path)
        assert len(wb.sheetnames) == 9
        assert wb.sheetnames == ["Reconciliation Summary", "Missing & Differences", "Taxable Value Missing", "Date Mismatches", "Invoice Mismatches", "Cancelled Vouchers", "Invoice No Remarks", "Not in Tally", "Notes"]
        
    def test_summary_sheet_layout(self, generated_excel_path):
        wb = load_workbook(generated_excel_path)
        ws = wb["Reconciliation Summary"]
        assert ws.cell(row=1, column=1).value == "#"
        assert ws.cell(row=1, column=12).value == "Status"
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None
        
        # Totals formulas
        totals_row = ws.max_row
        assert ws.cell(row=totals_row, column=1).value == "TOTALS"
        assert str(ws.cell(row=totals_row, column=8).value).startswith("=SUM(")
        assert ws["P6"].value == "Important"
        assert "do not determine ITC eligibility or GST compliance" in ws["Q6"].value

    def test_missing_and_differences_sheet(self, generated_excel_path):
        wb = load_workbook(generated_excel_path)
        ws = wb["Missing & Differences"]
        
        # Test Freeze Panes correctly at A3 to capture header
        assert ws.freeze_panes == "A3"
        
        # Test Tables (AutoFilter equivalent for multiple ranges)
        tables = list(ws.tables.values())
        assert len(tables) == 2, "Expected two Tables for Missing and Major Differences sections"
        assert any(t.displayName == "MissingInvoices" for t in tables)
        assert any(t.displayName == "MajorDifferences" for t in tables)

        # Check values
        assert ws.cell(row=1, column=1).value == "Invoices in Accounts but Missing in GSTR-2B"
        assert ws.cell(row=3, column=4).value == "INV002"
        
        # Tax differences are further down
        # Diff rows start from 7 (Row 5 is section header, row 6 is table header)
        assert ws.cell(row=5, column=1).value == "Major Differences"
        # Since diff table has CGST and SGST
        assert ws.cell(row=7, column=5).value == "CGST"
        assert ws.cell(row=7, column=8).value == D("2.00")
        assert ws.cell(row=8, column=5).value == "SGST"
        assert ws.cell(row=8, column=8).value == D("-2.00") # Signed diff

    def test_other_sheets(self, generated_excel_path):
        wb = load_workbook(generated_excel_path)
        
        ws3 = wb["Invoice No Remarks"]
        assert ws3.freeze_panes == "A2"
        assert ws3.cell(row=2, column=6).value == "Invoice Last 4"
        
        ws4 = wb["Not in Tally"]
        assert ws4.freeze_panes == "A2"
        assert ws4.cell(row=2, column=4).value == "NIT001"
        
        ws5 = wb["Notes"]
        assert ws5.freeze_panes == "A2"
        assert ws5.cell(row=2, column=5).value == "DN001"

    def test_formats_and_colors(self, generated_excel_path):
        wb = load_workbook(generated_excel_path)
        ws = wb["Reconciliation Summary"]
        
        # Decimal values
        val_cell = ws.cell(row=2, column=8)
        assert val_cell.number_format == '"₹"#,##0.00'
        
        # Date formats
        date_cell = ws.cell(row=2, column=7)
        assert date_cell.number_format == 'DD-MM-YYYY'
        
        # Color formats
        # Green for perfectly matched
        perfect_match_cell = ws.cell(row=2, column=12)
        assert perfect_match_cell.fill.start_color.rgb == "00C6EFCE"
        
        # Red for missing
        missing_cell = ws.cell(row=3, column=12)
        assert missing_cell.fill.start_color.rgb == "00FF9999"

    def test_invoice_mismatch_sheet_regression(self, tmp_path, invoice_mismatch_summary, mock_notes):
        path = tmp_path / "invoice-mismatch.xlsx"
        export_to_excel(str(path), invoice_mismatch_summary, mock_notes, "May 2025")
        ws = load_workbook(path)["Invoice Mismatches"]
        assert ws.cell(row=2, column=4).value == "SGST-5725"
        assert ws.cell(row=2, column=5).value == "SGST-5728"

    @pytest.mark.parametrize("consumed,expected_rows", [(False, 1), (True, 0)])
    def test_notes_export_deduplicates_and_excludes_consumed_candidates(
        self, tmp_path, mock_summary, consumed, expected_rows
    ):
        path = tmp_path / f"notes-{consumed}.xlsx"
        export_to_excel(
            str(path), mock_summary,
            _notes_with_compatibility_duplicate(consumed=consumed), "June 2025",
        )
        ws = load_workbook(path)["Notes"]
        exported = [
            row for row in ws.iter_rows(min_row=2, values_only=True)
            if row[5] == "CN-DUP-1"
        ]
        assert len(exported) == expected_rows


class TestJSONExporter:
    def test_json_structure(self, parsed_json):
        data = parsed_json
        assert "metadata" in data
        assert "counts" in data
        assert "reconciliation_summary" in data
        assert "invoice_mismatches" in data
        assert "missing_in_gstr" in data
        assert "major_differences" in data
        assert "invoice_number_remarks" in data
        assert "not_in_tally" in data
        assert "notes" in data
        assert "purchase_summary" in data
        assert "notes_summary" in data
        assert all(item["reconciliation_type"] == "PURCHASE" for item in data["reconciliation_summary"])
        assert all(item["reconciliation_type"] == "NOTE" for item in data["notes"])

    def test_json_metadata_and_counts(self, parsed_json):
        data = parsed_json
        assert data["metadata"]["reporting_period"] == "April 2025"
        assert data["counts"]["perfectly_matched"] == 2
        assert data["counts"]["missing_in_gstr_2b"] == 1
        assert data["counts"]["major_difference"] == 1
        assert data["counts"]["needs_review_ambiguous"] == 1
        assert data["counts"]["not_in_tally"] == 1
        assert data["counts"]["invoice_mismatches"] == 0
        assert data["invoice_mismatches"] == []
        assert data["metadata"]["reconciliation_scope"] == (
            "Reconciliation results reflect data matching between Accounts and GSTR records only. "
            "They do not determine ITC eligibility or GST compliance. "
            "Review eligibility separately before claiming ITC."
        )

    def test_invoice_mismatch_complete_filtered_view(self, invoice_mismatch_summary, mock_notes):
        data = json.loads(generate_json_report(invoice_mismatch_summary, mock_notes, "May 2025"))
        assert data["counts"]["invoice_mismatches"] == 1
        assert len(data["invoice_mismatches"]) == 1
        item = data["invoice_mismatches"][0]
        assert item["result_id"] == "result-mismatch-1"
        assert item["status"] == "Invoice Number Mismatch"
        assert item["invoice_mismatch"] is True
        assert item["tally_invoice_number"] == "SGST-5725"
        assert item["gstr_invoice_number"] == "SGST-5728"
        assert item["gstin"] == "06SYNTH1234A1Z5"
        assert item["tally_taxable_value"] == "1000.00"
        assert item["gstr_taxable_value"] == "1000.00"
        assert item["tally_cgst"] == item["gstr_cgst"] == "90.00"
        assert item["tally_sgst"] == item["gstr_sgst"] == "90.00"
        assert item["date_mismatch"] is True
        assert item["date_difference_days"] == 1
        assert item["evidence"] == "Exact GSTIN, date and financial evidence"
        assert not any(r["invoice_number_tally"] == "SGST-5725" for r in data["missing_in_gstr"])
        assert not any(r["invoice_number"] == "SGST-5728" for r in data["not_in_tally"])
        assert any(r["result_id"] == "result-mismatch-1" for r in data["reconciliation_summary"])

    def test_json_formatting(self, parsed_json):
        data = parsed_json
        
        # Decimal serialization as exact strings
        tax_diff_res = next(r for r in data["reconciliation_summary"] if r["status"] == "Major Difference")
        assert tax_diff_res["cgst_difference"] == "2.00"
        assert tax_diff_res["sgst_difference"] == "-2.00"
        
        # Enums readable
        assert tax_diff_res["match_method"] == "Exact Invoice"
        
        # Not in tally notes serialization
        not_in_tally_note = next(r for r in data["notes"] if r["status"] == "Not in Tally")
        assert not_in_tally_note["supplier_gst"] == "Supplier 6"
        assert not_in_tally_note["match_method"] == "No Match"

    @pytest.mark.parametrize("consumed,expected_rows", [(False, 1), (True, 0)])
    def test_notes_json_deduplicates_and_excludes_consumed_candidates(
        self, mock_summary, consumed, expected_rows
    ):
        data = json.loads(generate_json_report(
            mock_summary,
            _notes_with_compatibility_duplicate(consumed=consumed),
            "June 2025",
        ))
        exported = [
            item for item in data["notes"]
            if item.get("invoice_number_gst") == "CN-DUP-1"
        ]
        assert len(exported) == expected_rows


class TestSummaryGenerator:
    def test_summary_text(self, mock_summary, mock_notes):
        text = generate_text_summary(mock_summary, mock_notes)
        assert "Accounts invoices: 10" in text
        assert "Missing in GSTR-2B: 1" in text
        assert "--- Missing Invoices ---" in text
        assert "- Invoice INV002 from Supplier 2" in text
        assert "--- Major Differences ---" in text
        assert "- Invoice INV003 from Supplier 3" in text
