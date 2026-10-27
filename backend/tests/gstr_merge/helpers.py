import json
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

import openpyxl

from backend.app.services.gstr_merge.integrity_checker import IntegrityChecker
from backend.app.services.gstr_merge.processor import MergeContext, process_workbooks
from backend.app.services.gstr_merge.workbook_inspector import WorkbookInspector
from backend.app.services.gstr_merge.workbook_merger import WorkbookMerger
from backend.app.services.gstr_merge.workbook_validator import WorkbookValidator


GSTIN = "06AAAAA0000A1Z5"
MONTHS = [
    "April", "May", "June", "July", "August", "September", "October",
    "November", "December", "January", "February", "March",
]

HEADERS = {
    "B2B": ["GSTIN/UIN", "Invoice No.", "Invoice Date", "Taxable Value", "IGST", "CGST", "SGST", "Cess Amount"],
    "B2BA": ["GSTIN/UIN", "Original Invoice No.", "Original Invoice Date", "Revised Invoice No.", "Revised Invoice Date", "Taxable Value", "IGST"],
    "CDNR": ["GSTIN/UIN", "Note/Refund Number", "Note/Refund Date", "Taxable Value", "IGST", "CGST", "SGST", "Cess"],
    "CDNRA": ["GSTIN/UIN", "Original Note No.", "Original Note Date", "Revised Note Number", "Revised Note Date", "Taxable Value", "IGST"],
}


def row_for(category, index):
    if category == "B2B":
        return [GSTIN, f"INV-{index}", f"01-Apr-202{index % 10}", 100 + index, index, 2, 3, 0]
    if category == "B2BA":
        return [GSTIN, f"OLD-{index}", "01-Apr-2025", f"REV-{index}", "02-Apr-2025", 200 + index, 10]
    if category == "CDNR":
        return [GSTIN, f"NOTE-{index}", "03-Apr-2025", 50 + index, 1, 2, 3, 0]
    return [GSTIN, f"ON-{index}", "03-Apr-2025", f"RN-{index}", "04-Apr-2025", 60 + index, 2]


def make_book(path: Path, period=None, fy="2025-26", gstin=GSTIN, sheets=None, sheet_order=None):
    sheets = sheets or {"B2B": ("B2B", [row_for("B2B", 1)])}
    wb = openpyxl.Workbook()
    main = wb.active
    main.title = "Main"
    main.append(["GSTIN of the Tax Payer", gstin] if gstin else ["Name", "Test Taxpayer"])
    if fy:
        main.append(["Financial Year", fy])
    if period:
        main.append(["Months", period])
    names = sheet_order or list(sheets)
    for name in names:
        category, rows = sheets[name]
        ws = wb.create_sheet(name)
        ws.append(HEADERS[category])
        for row in rows:
            ws.append(row)
    wb.save(path)
    wb.close()
    return path


def make_period_books(tmp_path, mode, count=None, sheets_factory=None, metadata=True):
    tmp_path.mkdir(parents=True, exist_ok=True)
    count = count or (4 if mode == "quarterly" else 12)
    periods = (
        ["Apr-Jun 2025", "Jul-Sep 2025", "Oct-Dec 2025", "Jan-Mar 2026"]
        if mode == "quarterly"
        else [f"{month} {2025 if i < 9 else 2026}" for i, month in enumerate(MONTHS)]
    )
    paths = []
    for index in range(count):
        category_rows = (
            sheets_factory(index)
            if sheets_factory
            else {"B2B": ("B2B", [row_for("B2B", index + 1)])}
        )
        path = tmp_path / f"source-{index + 1}.xlsx"
        make_book(path, periods[index % len(periods)] if metadata else None, sheets=category_rows)
        paths.append(path)
    return paths


def files_payload(paths):
    """Compatibility helper retained for directly ported core assertions."""
    return paths


def add_sheet(paths, name, rows):
    for path in paths:
        wb = openpyxl.load_workbook(path)
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        wb.save(path)
        wb.close()


@dataclass
class CoreResult:
    status_code: int
    text: str
    content: bytes = b""
    summary: dict | None = None

    @property
    def headers(self):
        return {"X-Audit-Summary": json.dumps(self.summary)} if self.summary is not None else {}

    def json(self):
        return {"detail": self.text}


def run_core(paths, mode, tmp_path, **context_overrides):
    """Invoke the core directly; no HTTP route or application wiring is involved."""
    path_items = list(paths)
    mapping = {}
    filenames = {}
    for index, item in enumerate(path_items, 1):
        source_id = f"file_{index:02d}"
        if isinstance(item, tuple):
            filename = item[1][0]
            path = tmp_path / filename
        else:
            path = Path(item)
            filename = path.name
        mapping[source_id] = str(path)
        filenames[source_id] = filename
    output = tmp_path / f"core-result-{uuid4().hex}.xlsx"
    context = MergeContext(mode=mode, **context_overrides)
    try:
        summary = process_workbooks(mapping, filenames, context, output)
        return CoreResult(200, "", output.read_bytes(), summary)
    except (ValueError, OSError) as exc:
        return CoreResult(400, str(exc))
    finally:
        output.unlink(missing_ok=True)


def replace_with_official_gstr2b_read_me(path, tax_period, filename_gstin=GSTIN):
    wb = openpyxl.load_workbook(path)
    if "Main" in wb.sheetnames:
        wb.remove(wb["Main"])
    read_me = wb.create_sheet("Read me", 0)
    metadata = [
        ("Financial Year", "2025-26"),
        ("Tax Period", tax_period),
        ("GSTIN of taxpayer", filename_gstin),
        ("Legal Name", "Synthetic Legal Name"),
        ("Trade Name (if any)", "Synthetic Trade Name"),
        ("Date of Generation", "14/05/2025"),
    ]
    for row_number, (label, value) in enumerate(metadata, 4):
        read_me.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=2)
        read_me.cell(row_number, 1, label)
        read_me.cell(row_number, 3, value)
    for title in ("ITC Available", "ITC not available", "ITC Reversal", "ITC Rejected"):
        ws = wb.create_sheet(title)
        ws.append(["Description", "Amount"])
        ws.append([tax_period, 1])
    wb.save(path)
    wb.close()


def build_direct_merge(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    mapping = {f"s{i}": str(path) for i, path in enumerate(paths)}
    metadata = WorkbookInspector(mapping).inspect_all()
    assert WorkbookValidator(metadata, "quarterly").validate() == []
    output = str(tmp_path / "merged.xlsx")
    merger = WorkbookMerger(mapping, metadata, output)
    manifest = merger.merge()
    assert IntegrityChecker(manifest, output).verify() == []
    return output, manifest
