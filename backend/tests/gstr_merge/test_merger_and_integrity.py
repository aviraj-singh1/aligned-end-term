import json
import os
from pathlib import Path
import shutil
import tempfile
import zipfile

import openpyxl
import pytest

from backend.app.services.gstr_merge import integrity_checker as integrity_module
from backend.app.services.gstr_merge import source_manifest as source_manifest_module
from backend.app.services.gstr_merge.integrity_checker import IntegrityChecker
from backend.app.services.gstr_merge.package_safety import PackageSafetyError, validate_xlsx_package
from backend.app.services.gstr_merge.processor import MergeContext, ProcessingValidationError, process_workbooks
from backend.app.services.gstr_merge.workbook_inspector import WorkbookInspector
from backend.app.services.gstr_merge.workbook_merger import MergeValidationError, WorkbookMerger
from backend.app.services.gstr_merge.workbook_validator import WorkbookValidator
from backend.app.services.gstr_merge.source_manifest import (
    EXCEL_MAX_ROWS,
    SourceManifestBuilder,
    SourceValidationError,
    iter_spooled_rows,
    typed_row_token,
    validate_excel_row_limits,
)
from backend.tests.gstr_merge.helpers import (
    GSTIN,
    HEADERS,
    MONTHS,
    build_direct_merge,
    files_payload,
    make_book,
    make_period_books,
    replace_with_official_gstr2b_read_me,
    row_for,
    run_core,
)


@pytest.mark.parametrize("mode,count", [("quarterly", 4), ("monthly", 12)])
def test_full_core_modes_generate_reopenable_workbook(tmp_path, mode, count):
    paths = make_period_books(tmp_path, mode, count)
    response = run_core(paths, mode, tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["files_processed"] == count
    assert summary["rows_merged_per_category"]["B2B"] == count
    assert summary["integrity_status"] == "passed"
    output = tmp_path / "result.xlsx"
    output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=False)
    assert wb["B2B"].max_row == count + 1
    wb.close()


def test_optional_categories_aliases_order_case_and_column_mapping(tmp_path):
    periods = ["Apr-Jun 2025", "Jul-Sep 2025", "Oct-Dec 2025", "Jan-Mar 2026"]
    configurations = [
        {"B2B": ("B2B", [row_for("B2B", 1)]), "CDNR": ("CDNR", [row_for("CDNR", 1)])},
        {" notes ": ("CDNR", [row_for("CDNR", 2)]), " b2b ": ("B2B", [row_for("B2B", 2)])},
        {"B2BA": ("B2BA", [row_for("B2BA", 3)]), "b2B": ("B2B", [row_for("B2B", 3)])},
        {" cdnra ": ("CDNRA", [row_for("CDNRA", 4)]), "B2B": ("B2B", [row_for("B2B", 4)])},
    ]
    paths = []
    for index, sheets in enumerate(configurations):
        path = tmp_path / f"q{index + 1}.xlsx"
        make_book(path, periods[index], sheets=sheets, sheet_order=list(reversed(list(sheets))))
        paths.append(path)
    wb = openpyxl.load_workbook(paths[1])
    ws = wb[" b2b "]
    values = [cell.value for cell in ws[1]]
    order = [2, 0, 3, 1, 5, 4, 7, 6]
    ws.delete_rows(1, ws.max_row)
    ws.append([values[i] for i in order])
    source_row = row_for("B2B", 2)
    ws.append([source_row[i] for i in order])
    wb.save(paths[1]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["rows_merged_per_category"] == {"B2B": 4, "B2BA": 1, "CDNR": 2, "CDNRA": 1}
    output = tmp_path / "union.xlsx"; output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=True)
    assert {"B2B", "B2BA", "CDNR", "CDNRA"}.issubset(wb.sheetnames)
    assert wb["B2B"].cell(3, 2).value == "INV-2"
    wb.close()


def test_gstr2b_quarterly_detail_tabs_merge_and_period_summaries_are_omitted(tmp_path):
    periods = ["Apr-Jun 2025", "Jul-Sep 2025", "Oct-Dec 2025", "Jan-Mar 2026"]
    paths = []
    for index, period in enumerate(periods, 1):
        path = make_book(
            tmp_path / f"gstr2b-q{index}.xlsx",
            period,
            sheets={
                "B2B": ("B2B", [[GSTIN, f"INV-{index}", "01-Apr-2025", 100 + index, 18, 0, 0, 0]]),
            },
        )
        wb = openpyxl.load_workbook(path)
        b2b = wb["B2B"]
        b2b.delete_rows(1, b2b.max_row)
        b2b.append([
            "GSTIN of supplier", "Invoice number", "Invoice Date", "Taxable Value (₹)",
            "Integrated Tax (₹)", "Central Tax (₹)", "State/UT Tax (₹)", "Cess (₹)",
        ])
        b2b.append([GSTIN, f"INV-{index}", "01-Apr-2025", 100 + index, 18, 0, 0, 0])

        cdnr = wb.create_sheet("B2B-CDNR")
        cdnr.append([
            "GSTIN of supplier", "Note number", "Note date", "Taxable Value (₹)",
            "Integrated Tax (₹)",
        ])
        cdnr.append([GSTIN, "TOTAL" if index == 1 else f"CN-{index}", "02-Apr-2025", 10 + index, 2])

        impg = wb.create_sheet("IMPGSEZ")
        impg.append([
            "GSTIN of supplier", "Bill of Entry No.", "Bill of Entry Date",
            "Taxable Value (₹)", "Integrated Tax (₹)",
        ])
        impg.append([GSTIN, f"BOE-{index}", "03-Apr-2025", 200 + index, 36])

        for title in ("ITC Available", "ITC not available", "ITC Reversal", "ITC Rejected"):
            summary = wb.create_sheet(title)
            summary.append(["Description", "Integrated Tax (₹)"])
            summary.append([f"Quarter {index}", index * 100])
        read_me = wb.create_sheet("Read me")
        read_me.append(["Section", "Instruction"])
        read_me.append(["GSTR-2B", "Portal generated statement"])
        wb.save(path)
        wb.close()
        paths.append(path)

    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["rows_merged_per_category"] == {
        "B2B": 4, "IMPG SEZ": 4, "B2B-CDNR": 4,
    }
    assert summary["validation_warnings"] == []
    assert len(summary["informational_notes"]) == 4

    output = tmp_path / "gstr2b-annual.xlsx"
    output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=True)
    assert wb.sheetnames[0] == "Annual Summary"
    assert "Read me" in wb.sheetnames
    assert "ITC Available" not in wb.sheetnames
    assert wb["B2B-CDNR"].cell(2, 2).value == "TOTAL"
    assert wb["B2B"].max_row == 5
    wb.close()


def test_gstr2b_stacked_headers_and_repeated_header_blocks_merge(tmp_path):
    periods = ["Apr-Jun 2025", "Jul-Sep 2025", "Oct-Dec 2025", "Jan-Mar 2026"]
    paths = []
    b2b_top = [
        "GSTIN of supplier", "Invoice number", "Invoice Date", "Invoice Value (₹)",
        "Tax Amount", None, None,
    ]
    b2b_bottom = [None, None, None, None, "Taxable Value (₹)", "Integrated Tax(₹)", "Cess(₹)"]
    b2ba_top = [
        "Original Details", None, "Revised Details", None, None,
        "Tax Amount", None, "Amount declared by taxpayer for ITC reduction", None,
    ]
    b2ba_bottom = [
        "Invoice number", "Invoice Date", "GSTIN of supplier", "Invoice number", "Invoice Date",
        "Taxable Value (₹)", "Integrated Tax(₹)", "Integrated Tax(₹)", "Cess(₹)",
    ]
    impg_top = ["ICEGATE Reference Date", "Port Code", "Bill of Entry Details", None, "Amount of tax (₹)", None]
    impg_bottom = [None, None, "Number", "Date", "Taxable Value", "Integrated Tax(₹)"]

    for index, period in enumerate(periods, 1):
        path = make_book(tmp_path / f"stacked-q{index}.xlsx", period)
        wb = openpyxl.load_workbook(path)
        b2b = wb["B2B"]
        b2b.delete_rows(1, b2b.max_row)
        b2b.append(b2b_top)
        b2b.append(b2b_bottom)
        b2b.append([GSTIN, f"INV-{index}", "01-Apr-2025", 118, 100, 18, 0])
        if index == 1:
            b2b.append(b2b_top)
            b2b.append(b2b_bottom)
            b2b.append([GSTIN, "SUMMARY", "02-Apr-2025", 236, 200, 36, 0])

        b2ba = wb.create_sheet("B2BA")
        b2ba.append(b2ba_top)
        b2ba.append(b2ba_bottom)
        b2ba.append(["OLD-1", "01-Apr-2025", GSTIN, f"NEW-{index}", "02-Apr-2025", 50, 9, 1, 0])

        impg = wb.create_sheet("IMPG")
        impg.append(impg_top)
        impg.append(impg_bottom)
        impg.append(["03-Apr-2025", "INDEL4", f"BOE-{index}", "03-Apr-2025", 300, 54])
        wb.save(path)
        wb.close()
        paths.append(path)

    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["rows_merged_per_category"] == {"B2B": 5, "B2BA": 4, "IMPG": 4}
    output = tmp_path / "stacked-annual.xlsx"
    output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=True)
    assert "SUMMARY" in [wb["B2B"].cell(row, 2).value for row in range(3, 8)]
    assert wb["IMPG"].cell(3, 3).value == "BOE-1"
    assert wb["B2BA"].cell(3, 8).value == 1
    wb.close()


def test_empty_excel_strings_are_canonicalized_before_integrity_hashing(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    for path in paths:
        wb = openpyxl.load_workbook(path)
        wb["B2B"].cell(2, 8).value = ""
        wb.save(path)
        wb.close()

    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    assert json.loads(response.headers["X-Audit-Summary"])["integrity_status"] == "passed"


@pytest.mark.parametrize("column,field", [(4, "taxable_value"), (5, "igst"), (6, "cgst"), (7, "sgst"), (8, "cess")])
def test_integrity_detects_financial_mutation(tmp_path, column, field):
    output, manifest = build_direct_merge(tmp_path)
    wb = openpyxl.load_workbook(output); wb["B2B"].cell(2, column).value += 1; wb.save(output); wb.close()
    errors = IntegrityChecker(manifest, output).verify()
    assert any(field in error for error in errors)


@pytest.mark.parametrize("mutation", ["loss", "duplicate"])
def test_integrity_detects_row_multiplicity_changes(tmp_path, mutation):
    output, manifest = build_direct_merge(tmp_path)
    wb = openpyxl.load_workbook(output); ws = wb["B2B"]
    if mutation == "loss": ws.delete_rows(2)
    else: ws.append([cell.value for cell in ws[2]])
    wb.save(output); wb.close()
    errors = IntegrityChecker(manifest, output).verify()
    assert any("row" in error.lower() for error in errors)


def test_zero_blank_and_legitimate_duplicate_multiplicity(tmp_path):
    shared = row_for("B2B", 9)
    def factory(index):
        row = list(shared)
        row[4] = 0 if index == 0 else None
        return {"B2B": ("B2B", [row])}
    paths = make_period_books(tmp_path, "quarterly", sheets_factory=factory)
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["rows_merged_per_category"]["B2B"] == 4
    assert summary["duplicate_warnings"][0]["duplicate_count"] == 4
    output = tmp_path / "zero-blank.xlsx"; output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=True)
    values = [wb["B2B"].cell(row, 5).value for row in range(2, 6)]
    assert values == [0, None, None, None]
    wb.close()


def test_likely_note_duplicates_are_warned_not_removed(tmp_path):
    shared = row_for("CDNR", 7)
    paths = make_period_books(tmp_path, "quarterly", sheets_factory=lambda _: {"CDNR": ("CDNR", [shared])})
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["rows_merged_per_category"]["CDNR"] == 4
    assert summary["duplicate_warnings"][0]["category"] == "CDNR"


def test_unavailable_metadata_warns_without_fabrication(tmp_path):
    paths = make_period_books(tmp_path, "quarterly", metadata=False)
    for path in paths:
        wb = openpyxl.load_workbook(path); wb.remove(wb["Main"]); wb.save(path); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["gstin_validation"]["status"] == "unavailable"
    assert len(summary["period_warnings"]) == 4
    assert len(summary["validation_warnings"]) == 12


def test_integrity_detects_identifier_mutation_and_total_identifier_survives(tmp_path):
    paths = make_period_books(
        tmp_path,
        "quarterly",
        sheets_factory=lambda index: {"B2B": ("B2B", [[
            "06AAAAA0000A1Z5", "TOTAL" if index == 0 else f"INV-{index}",
            "01-Apr-2025", 100, 0, 9, 9, 0,
        ]])},
    )
    mapping = {f"s{i}": str(path) for i, path in enumerate(paths)}
    metadata = WorkbookInspector(mapping).inspect_all()
    manifest = SourceManifestBuilder(mapping, metadata).build()
    output = str(tmp_path / "total-integrity.xlsx")
    WorkbookMerger(mapping, metadata, output, manifest).merge()
    assert IntegrityChecker(manifest, output).verify() == []
    wb = openpyxl.load_workbook(output); assert wb["B2B"].cell(2, 2).value == "TOTAL"
    wb["B2B"].cell(2, 2).value = "ALTERED"; wb.save(output); wb.close()
    assert any("multiplicity/hash" in error for error in IntegrityChecker(manifest, output).verify())


def test_integrity_reopens_output_read_only(tmp_path, monkeypatch):
    output, manifest = build_direct_merge(tmp_path)
    assert all(item["spool_path"] is None for item in manifest["categories"].values())
    calls = []
    original = integrity_module.openpyxl.load_workbook

    def tracked_load(path, *args, **kwargs):
        calls.append(kwargs.copy())
        return original(path, *args, **kwargs)

    monkeypatch.setattr(integrity_module.openpyxl, "load_workbook", tracked_load)
    assert IntegrityChecker(manifest, output).verify() == []
    assert calls and calls[-1].get("read_only") is True


def test_streamed_output_preserves_date_format_and_merged_static_cells(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    for path in paths:
        wb = openpyxl.load_workbook(path)
        wb["Main"].merge_cells("C4:D4")
        wb["Main"]["C4"] = "Template note"
        wb["B2B"]["C2"] = openpyxl.utils.datetime.from_ISO8601("2025-04-01")
        wb["B2B"]["C2"].number_format = "dd-mmm-yyyy"
        wb.save(path)
        wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    output = tmp_path / "formatted.xlsx"
    output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=False)
    assert "C4:D4" in {str(item) for item in wb["Main"].merged_cells.ranges}
    assert wb["Main"]["C4"].value == "Template note"
    assert wb["B2B"]["C2"].number_format == "dd-mmm-yyyy"
    wb.close()


@pytest.mark.parametrize("mode,count", [("quarterly", 4), ("monthly", 12)])
def test_direct_core_smoke_verifies_rows_multiplicity_and_financial_totals(tmp_path, mode, count):
    def rows(index):
        return {"B2B": ("B2B", [[
            GSTIN,
            f"SYN-{index + 1}",
            "01-Apr-2025",
            100 + index,
            10,
            5,
            5,
            1,
        ]])}

    paths = make_period_books(tmp_path, mode, sheets_factory=rows)
    response = run_core(paths, mode, tmp_path)
    assert response.status_code == 200, response.text
    assert response.summary["integrity_status"] == "passed"
    assert response.summary["rows_merged_per_category"] == {"B2B": count}

    output = tmp_path / f"{mode}-smoke.xlsx"
    output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=False, read_only=True, keep_links=False)
    assert wb.sheetnames[0] == "Annual Summary"
    assert "B2B" in wb.sheetnames
    rows = list(wb["B2B"].iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(rows) == count
    assert len({row[1] for row in rows}) == count
    assert sum(row[3] for row in rows) == sum(100 + index for index in range(count))
    assert sum(row[4] for row in rows) == 10 * count
    assert sum(row[5] for row in rows) == 5 * count
    assert sum(row[6] for row in rows) == 5 * count
    assert sum(row[7] for row in rows) == count
