import json
import os
from pathlib import Path
import shutil
import tempfile
import zipfile

import openpyxl
import pytest

from backend.app.services.gstr_merge import integrity_checker as integrity_module
from backend.app.services.gstr_merge import source_manifest as source_manifest_module
from backend.app.services.gstr_merge.integrity_checker import IntegrityChecker
from backend.app.services.gstr_merge.package_safety import PackageSafetyError, validate_xlsx_package
from backend.app.services.gstr_merge.processor import MergeContext, ProcessingValidationError, process_workbooks
from backend.app.services.gstr_merge.workbook_inspector import WorkbookInspector
from backend.app.services.gstr_merge.workbook_merger import MergeValidationError, WorkbookMerger
from backend.app.services.gstr_merge.workbook_validator import WorkbookValidator
from backend.app.services.gstr_merge.source_manifest import (
    EXCEL_MAX_ROWS,
    SourceManifestBuilder,
    SourceValidationError,
    iter_spooled_rows,
    typed_row_token,
    validate_excel_row_limits,
)
from backend.tests.gstr_merge.helpers import (
    GSTIN,
    HEADERS,
    MONTHS,
    add_sheet,
    build_direct_merge,
    files_payload,
    make_book,
    make_period_books,
    replace_with_official_gstr2b_read_me,
    row_for,
    run_core,
)


@pytest.mark.parametrize("alias", ["CDNR", "Notes", "Note", "Credit Debit Notes"])
def test_cdnr_name_alias_requires_compatible_schema(tmp_path, alias):
    path = make_book(tmp_path / "alias.xlsx", sheets={alias: ("CDNR", [row_for("CDNR", 1)])})
    metadata = WorkbookInspector({"one": str(path)}).inspect_all()
    assert "cdnr" in metadata["one"]["categories"]


def test_gstr2b_currency_units_are_accepted_without_weakening_header_identity():
    from backend.app.services.gstr_merge.gst_schema import CATEGORY_BY_KEY, canonical_header, category_candidates

    assert canonical_header("Taxable Value (₹)") == "taxable_value"
    assert canonical_header("Invoice Value (Rs.)") == "invoice_value"
    assert canonical_header("Integrated Tax (INR)") == "igst"
    assert canonical_header("Rate (%)").startswith("raw:")
    assert canonical_header("Document (Amendment)").startswith("raw:")
    for sheet_name in (
        "B2B-CDNR", "B2B-CDNRA", "ECO", "ECOA", "ISDA", "IMPGA", "IMPGSEZ", "IMPGSEZA",
        "B2B (ITC Reversal)", "B2BA (ITC Reversal)", "B2B-DNR", "B2B-DNRA",
        "B2B(Rejected)", "B2BA(Rejected)", "B2B-CDNR(Rejected)", "B2B-CDNRA(Rejected)",
        "ECO(Rejected)", "ECOA(Rejected)", "ISD(Rejected)", "ISDA(Rejected)",
    ):
        assert len(category_candidates(sheet_name)) == 1

    spec = CATEGORY_BY_KEY["b2b"]
    portal_header = [
        "GSTIN of supplier", "Invoice number", "Invoice Date",
        "Taxable Value (₹)", "Integrated Tax (₹)",
    ]
    assert SourceManifestBuilder._is_repeated_header(portal_header, spec)
    assert not SourceManifestBuilder._is_repeated_header(
        [GSTIN, "Taxable Value (₹)", "01-Apr-2025", 100, 18], spec
    )


def test_amendment_sheet_accepts_portal_current_document_labels_in_context(tmp_path):
    from backend.app.services.gstr_merge.gst_schema import CATEGORY_BY_KEY

    path = make_book(tmp_path / "portal-b2ba.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["B2B"]
    ws.title = "B2BA"
    ws.delete_rows(1, ws.max_row)
    header = [
        "GSTIN of supplier", "Original Invoice number", "Original Invoice date",
        "Invoice number", "Invoice date", "Taxable Value (₹)", "Integrated Tax (₹)",
    ]
    ws.append(header)
    ws.append([GSTIN, "OLD-1", "01-Apr-2025", "NEW-1", "02-Apr-2025", 100, 18])
    ws.append(header)
    wb.save(path)
    wb.close()

    metadata = WorkbookInspector({"one": str(path)}).inspect_all()
    assert metadata["one"]["errors"] == []
    info = metadata["one"]["categories"]["b2ba"]
    assert info["column_map"]["revised_invoice_number"] == 3
    assert SourceManifestBuilder._is_repeated_header(header, CATEGORY_BY_KEY["b2ba"])


def test_official_gstr2b_read_me_metadata_overrides_quarter_filename_and_is_rejected_in_quarterly_mode(tmp_path):
    months = ["April", "July", "October", "January"]
    paths = []
    for index, month in enumerate(months, 1):
        path = make_book(tmp_path / f"synthetic_Q{index}.xlsx", sheets={"B2B": ("B2B", [row_for("B2B", index)])})
        replace_with_official_gstr2b_read_me(path, month)
        paths.append(path)

    metadata = WorkbookInspector(
        {f"s{index}": str(path) for index, path in enumerate(paths)}
    ).inspect_all()
    assert [item["period"]["code"] for item in metadata.values()] == ["M04", "M07", "M10", "M01"]
    assert all(item["period"]["source"] == "workbook" for item in metadata.values())
    assert all(item["taxpayer_gstin"] == GSTIN for item in metadata.values())
    assert all(item["financial_year"] == "2025-26" for item in metadata.values())
    assert all(item["legal_name"] == "Synthetic Legal Name" for item in metadata.values())
    assert all(item["trade_name"] == "Synthetic Trade Name" for item in metadata.values())

    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    detail = response.json()["detail"]
    assert "Quarterly mode received monthly GSTR-2B statements" in detail
    assert "April, July, October, January" in detail
    assert "upload all 12 statements" in detail


def test_stacked_header_detection_stops_before_header_like_business_values(tmp_path):
    path = make_book(tmp_path / "stacked-data.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["B2B"]
    ws.delete_rows(1, ws.max_row)
    ws.append(["GSTIN of supplier", "Trade/Legal name", "Invoice Details", None, None, "Tax Amount"])
    ws.append([None, None, "Invoice number", "Invoice Date", "Taxable Value (₹)", "Integrated Tax(₹)"])
    ws.append([GSTIN, "Invoice Date", "SUMMARY", "01-Apr-2025", 100, 18])
    ws.append([GSTIN, "Taxable Value", "TOTAL", "02-Apr-2025", 200, 36])
    wb.save(path)
    wb.close()

    metadata = WorkbookInspector({"q1": str(path)}).inspect_all()["q1"]
    assert metadata["errors"] == []
    assert metadata["categories"]["b2b"]["header_row_idx"] == 2


def test_amendment_stacked_headers_keep_duplicate_raw_fields_by_context(tmp_path):
    path = make_book(tmp_path / "amendments.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("B2B-CDNRA")
    ws.append(["Original Details", None, None, "Revised Details", None, None, None, None, None, None])
    ws.append([
        "Note type", "Note number", "Note date", "GSTIN of supplier", "Trade/Legal name",
        "Credit note/Debit note details", None, None, None, "Taxable Value (₹)",
    ])
    ws.append([None, None, None, None, None, "Note number", "Note type", "Note Supply type", "Note date", None])
    wb.save(path)
    wb.close()

    metadata = WorkbookInspector({"q1": str(path)}).inspect_all()["q1"]
    assert metadata["errors"] == []
    mapping = metadata["categories"]["gstr2b_b2b_cdnra"]["column_map"]
    assert mapping["raw:original:note type"] == 0
    assert mapping["raw:revised:note type"] == 6


def test_complete_top_header_includes_following_tax_subheader(tmp_path):
    path = make_book(tmp_path / "isd.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("ISD")
    ws.append([
        "GSTIN of ISD", "Trade/Legal name", "ISD Document type", "ISD Document number",
        "ISD Document date", "Original Invoice Number", "Original invoice date",
        "Input tax distribution by ISD", None, None, None,
    ])
    ws.append([None, None, None, None, None, None, None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"])
    wb.save(path)
    wb.close()

    metadata = WorkbookInspector({"q1": str(path)}).inspect_all()["q1"]
    assert metadata["errors"] == []
    info = metadata["categories"]["isd"]
    assert info["header_row_idx"] == 2
    assert info["column_map"]["igst"] == 7


def test_empty_portal_dnra_template_does_not_relax_financial_row_validation(tmp_path):
    path = make_book(tmp_path / "empty-dnra.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet("B2B-DNRA")
    ws.append(["Original Details", None, None, "Revised Details", None, None, None, None, None, None, None, None, None, None])
    ws.append([
        "Note type", "Note number", "Note date", "GSTIN of supplier", "Trade/Legal name",
        "Debit note details", None, None, None, None, "Place of supply", "Supply Attract Reverse Charge",
        "Rate(%)", "Taxable Value (₹)",
    ])
    ws.append([None, None, None, None, None, "Note number", "Note type", "Note Supply type", "Note date", "Note Value (₹)", None, None, None, "Integrated Tax(₹)"])
    wb.save(path)
    wb.close()

    metadata = WorkbookInspector({"q1": str(path)}).inspect_all()["q1"]
    assert metadata["errors"] == []
    assert "gstr2b_b2b_dnra" in metadata["categories"]

    wb = openpyxl.load_workbook(path)
    wb["B2B-DNRA"].append(["Debit Note", "OLD-1", "01-Apr-2025", GSTIN, "Supplier", "NEW-1", "Debit Note", "Regular", "02-Apr-2025", 100])
    wb.save(path)
    wb.close()
    errors = WorkbookInspector({"q1": str(path)}).inspect_all()["q1"]["errors"]
    assert any("B2B-DNRA" in error and "could not find" in error for error in errors)


def test_missing_duplicate_and_unmappable_headers_fail(tmp_path):
    missing = make_book(tmp_path / "missing.xlsx")
    wb = openpyxl.load_workbook(missing); wb["B2B"].cell(1, 2).value = "Reference"; wb.save(missing); wb.close()
    assert "could not find" in " ".join(WorkbookInspector({"x": str(missing)}).inspect_all()["x"]["errors"])

    duplicate = make_book(tmp_path / "duplicate.xlsx")
    wb = openpyxl.load_workbook(duplicate); wb["B2B"].cell(1, 1).value = "Invoice Number"; wb.save(duplicate); wb.close()
    assert "duplicate/ambiguous" in " ".join(WorkbookInspector({"x": str(duplicate)}).inspect_all()["x"]["errors"])

    paths = make_period_books(tmp_path / "mapped", "quarterly")
    wb = openpyxl.load_workbook(paths[1]); ws = wb["B2B"]; ws.cell(1, 9).value = "Extra Field"; ws.cell(2, 9).value = "must-not-drop"; wb.save(paths[1]); wb.close()
    metadata = WorkbookInspector({f"s{i}": str(path) for i, path in enumerate(paths)}).inspect_all()
    with pytest.raises(MergeValidationError, match="cannot be mapped"):
        WorkbookMerger({f"s{i}": str(path) for i, path in enumerate(paths)}, metadata, str(tmp_path / "bad.xlsx")).merge()


def test_repeated_header_detection_keeps_full_reordered_header_semantics():
    from backend.app.services.gstr_merge.gst_schema import CATEGORY_BY_KEY

    spec = CATEGORY_BY_KEY["b2b"]
    assert SourceManifestBuilder._is_repeated_header(HEADERS["B2B"], spec)
    assert SourceManifestBuilder._is_repeated_header(
        ["Taxable Value", "Invoice Date", "unmapped heading", "Invoice No."], spec
    )
    assert not SourceManifestBuilder._is_repeated_header(
        ["06AAAAA0000A1Z5", "Invoice No.", "01-Apr-2025", 100, 0, 0, 0, 0], spec
    )


def test_identical_unknown_non_tabular_sheet_is_retained(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    add_sheet(paths, "Read Me", [["Generator version 1"]])
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    output = tmp_path / "unknown.xlsx"; output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output)
    assert wb["Read Me"]["A1"].value == "Generator version 1"
    wb.close()


def test_unknown_differing_sheet_is_rejected(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    add_sheet(paths, "Read Me", [["Generator version 1"]])
    wb = openpyxl.load_workbook(paths[-1]); wb["Read Me"]["A1"] = "Generator version 2"; wb.save(paths[-1]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert "differs across uploaded files" in response.json()["detail"]


def test_unknown_populated_table_is_rejected(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    add_sheet(paths, "Purchase Register", [["Invoice", "Amount"], ["P-1", 100]])
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert "Unsupported populated table" in response.json()["detail"]


def test_approved_static_and_metadata_sheets_are_retained(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    add_sheet(paths, "Party", [["Name", "Sanitized taxpayer"]])
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    output = tmp_path / "approved.xlsx"; output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output)
    assert {"Main", "Party"}.issubset(wb.sheetnames)
    wb.close()


def test_metadata_verification_states(tmp_path):
    full = make_period_books(tmp_path / "full", "quarterly")
    response = run_core(full, "quarterly", tmp_path)
    verification = json.loads(response.headers["X-Audit-Summary"])["metadata_verification"]
    assert {value["status"] for value in verification.values()} == {"VERIFIED"}

    partial = make_period_books(tmp_path / "partial", "quarterly")
    wb = openpyxl.load_workbook(partial[0]); wb["Main"]["B3"] = None; wb.save(partial[0]); wb.close()
    response = run_core(partial, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    verification = json.loads(response.headers["X-Audit-Summary"])["metadata_verification"]
    assert verification["period"]["status"] == "PARTIALLY VERIFIED"

    unavailable = make_period_books(tmp_path / "none", "quarterly", metadata=False)
    response = run_core(unavailable, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    verification = json.loads(response.headers["X-Audit-Summary"])["metadata_verification"]
    assert verification["period"]["status"] == "NOT VERIFIABLE"
