import json
import os
from pathlib import Path
import shutil
import tempfile
import zipfile

import openpyxl
import pytest

from backend.app.services.gstr_merge import integrity_checker as integrity_module
from backend.app.services.gstr_merge import package_safety as package_safety_module
from backend.app.services.gstr_merge import processor as processor_module
from backend.app.services.gstr_merge import source_manifest as source_manifest_module
from backend.app.services.gstr_merge.integrity_checker import IntegrityChecker
from backend.app.services.gstr_merge.package_safety import PackageSafetyError, validate_xlsx_package
from backend.app.services.gstr_merge.processor import MergeContext, ProcessingValidationError, process_workbooks
from backend.app.services.gstr_merge.workbook_inspector import WorkbookInspector
from backend.app.services.gstr_merge.workbook_merger import MergeValidationError, WorkbookMerger
from backend.app.services.gstr_merge.workbook_validator import WorkbookValidator
from backend.app.services.gstr_merge.source_manifest import (
    EXCEL_MAX_ROWS,
    SourceManifestBuilder,
    SourceValidationError,
    iter_spooled_rows,
    typed_row_token,
    validate_excel_row_limits,
)
from backend.tests.gstr_merge.helpers import (
    GSTIN,
    HEADERS,
    MONTHS,
    build_direct_merge,
    files_payload,
    make_book,
    make_period_books,
    replace_with_official_gstr2b_read_me,
    row_for,
    run_core,
)


@pytest.mark.parametrize("mode,count", [("quarterly", 3), ("quarterly", 5), ("monthly", 11), ("monthly", 13)])
def test_wrong_upload_counts_are_rejected(mode, count, tmp_path):
    payload = [("files", (f"f{i}.xlsx", b"x", "application/octet-stream")) for i in range(count)]
    response = run_core(payload, mode, tmp_path)
    assert response.status_code == 400
    assert "requires exactly" in response.json()["detail"]


def test_monthly_mode_requires_all_fiscal_year_months_without_duplicates(tmp_path):
    paths = make_period_books(tmp_path, "monthly")
    wb = openpyxl.load_workbook(paths[-1])
    wb["Main"]["B3"] = "February 2026"
    wb.save(paths[-1])
    wb.close()
    response = run_core(paths, "monthly", tmp_path)
    assert response.status_code == 400
    assert "Duplicate return period M02" in response.json()["detail"]
    assert "missing: M03" in response.json()["detail"]


def test_portal_na_is_only_allowed_for_optional_itc_reduction_amounts():
    assert SourceManifestBuilder._decimal(
        "NA", "source.xlsx", "B2B-CDNR", 7, "itc_reduction_igst"
    ) == 0
    with pytest.raises(SourceValidationError, match="expected numeric value"):
        SourceManifestBuilder._decimal("NA", "source.xlsx", "B2B-CDNR", 7, "taxable_value")
    with pytest.raises(SourceValidationError, match="expected numeric value"):
        SourceManifestBuilder._decimal("TOTAL", "source.xlsx", "B2B-CDNR", 7, "itc_reduction_igst")


def test_typed_row_hash_ignores_equal_string_object_sharing():
    shared = "supplier-name-with-enough-length-to-avoid-interning"
    separate_left = "".join(["supplier-name-with-", "enough-length-to-avoid-interning"])
    separate_right = "".join(["supplier-name-with-enough-", "length-to-avoid-interning"])
    assert shared == separate_left == separate_right
    assert separate_left is not separate_right
    assert typed_row_token((shared, shared)) == typed_row_token((separate_left, separate_right))


def test_duplicate_source_file_is_blocked(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    paths[1] = paths[0]
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert "Duplicate source files" in response.json()["detail"]


def test_taxpayer_period_and_year_validation(tmp_path):
    paths = make_period_books(tmp_path / "gstin", "quarterly")
    wb = openpyxl.load_workbook(paths[3]); wb["Main"]["B1"] = "27AAAAA0000A1Z5"; wb.save(paths[3]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert "Conflicting taxpayer GSTIN" in response.json()["detail"]

    paths = make_period_books(tmp_path / "period", "quarterly")
    wb = openpyxl.load_workbook(paths[1]); wb["Main"]["B3"] = "Apr-Jun 2025"; wb.save(paths[1]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert "Duplicate return period Q1" in response.json()["detail"]

    paths = make_period_books(tmp_path / "fy", "quarterly")
    wb = openpyxl.load_workbook(paths[2]); wb["Main"]["B2"] = "2026-27"; wb.save(paths[2]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert "Inconsistent financial years" in response.json()["detail"]


@pytest.mark.parametrize("kind", ["corrupt", "empty", "unrelated", "malformed"])
def test_invalid_workbooks_are_rejected(tmp_path, kind):
    paths = make_period_books(tmp_path, "quarterly")
    target = paths[0]
    if kind == "corrupt": target.write_bytes(b"not an xlsx")
    else:
        wb = openpyxl.Workbook(); ws = wb.active
        if kind == "unrelated": ws.append(["Name", "Value"]); ws.append(["x", 1])
        elif kind == "malformed": ws.title = "B2B"; ws.append(["Invoice Number", "Taxable Value"])
        wb.save(target); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400


def test_formula_in_business_data_is_rejected(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    wb = openpyxl.load_workbook(paths[0]); wb["B2B"].cell(2, 4).value = "=50+50"; wb.save(paths[0]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert "formulas in business data" in response.json()["detail"]


def test_repeated_headers_and_total_rows_are_not_merged(tmp_path):
    def factory(index):
        rows = [row_for("B2B", index + 1)]
        if index == 0:
            rows.extend([HEADERS["B2B"], ["Total", None, None, 9999, 999, 999, 999, 999]])
        return {"B2B": ("B2B", rows)}
    paths = make_period_books(tmp_path, "quarterly", sheets_factory=factory)
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["rows_merged_per_category"]["B2B"] == 4


@pytest.mark.parametrize(
    "identifier",
    ["TOTAL", "SUBTOTAL", "SUMMARY", "Invoice No.", "Taxable Value", "GSTIN/UIN"],
)
def test_header_and_footer_like_invoice_identifiers_are_preserved(tmp_path, identifier):
    def factory(index):
        row = row_for("B2B", index + 1)
        if index == 0:
            row[1] = identifier
        return {"B2B": ("B2B", [row])}

    paths = make_period_books(tmp_path, "quarterly", sheets_factory=factory)
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    output = tmp_path / "identifiers.xlsx"; output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=False)
    assert identifier in [wb["B2B"].cell(row, 2).value for row in range(2, 6)]
    wb.close()


def test_total_note_identifier_is_preserved(tmp_path):
    def factory(index):
        row = row_for("CDNR", index + 1)
        if index == 0:
            row[1] = "TOTAL"
        return {"CDNR": ("CDNR", [row])}

    paths = make_period_books(tmp_path, "quarterly", sheets_factory=factory)
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    output = tmp_path / "notes.xlsx"; output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output)
    assert wb["CDNR"].cell(2, 2).value == "TOTAL"
    wb.close()


@pytest.mark.parametrize("label", ["Total", "Total Amount"])
def test_structurally_identifiable_footer_is_excluded(tmp_path, label):
    def factory(index):
        rows = [row_for("B2B", index + 1)]
        if index == 0:
            rows.append([label, None, None, 9999, 999, 999, 999, 999])
        return {"B2B": ("B2B", rows)}

    paths = make_period_books(tmp_path, "quarterly", sheets_factory=factory)
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    summary = json.loads(response.headers["X-Audit-Summary"])
    assert summary["rows_merged_per_category"]["B2B"] == 4


def test_ambiguous_footer_word_is_not_classified_as_footer():
    from backend.app.services.gstr_merge.gst_schema import CATEGORY_BY_KEY
    values = ["06AAAAA0000A1Z5", "TOTAL", "01-Apr-2025", 100]
    mapping = {"gstin": 0, "invoice_number": 1, "invoice_date": 2, "taxable_value": 3}
    assert not SourceManifestBuilder._is_structural_footer(values, mapping, CATEGORY_BY_KEY["b2b"])


@pytest.mark.parametrize(
    "column,value,expected",
    [
        (1, "INVALID", "invalid GSTIN"),
        (2, None, "Invoice Number: value is required"),
        (3, "04/31/2025", "invalid or ambiguous date"),
        (4, "ABC", "expected numeric value"),
    ],
)
def test_row_level_validation_errors(tmp_path, column, value, expected):
    paths = make_period_books(tmp_path, "quarterly")
    wb = openpyxl.load_workbook(paths[0]); wb["B2B"].cell(2, column).value = value; wb.save(paths[0]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert "source-1.xlsx" in response.json()["detail"]
    assert "B2B" in response.json()["detail"]
    assert "row 2" in response.json()["detail"]
    assert expected in response.json()["detail"]


@pytest.mark.parametrize(
    "row_count,has_error",
    [
        (EXCEL_MAX_ROWS - 2, False),
        (EXCEL_MAX_ROWS - 1, False),
        (EXCEL_MAX_ROWS, True),
    ],
)
def test_excel_row_limit_preflight_without_large_workbook(row_count, has_error):
    manifest = {"categories": {"b2b": {"header_row_idx": 1, "row_count": row_count}}}
    assert bool(validate_excel_row_limits(manifest)) is has_error


def test_internal_static_formula_is_preserved_and_external_formula_rejected(tmp_path):
    paths = make_period_books(tmp_path / "internal", "quarterly")
    for path in paths:
        wb = openpyxl.load_workbook(path); wb["Main"]["C1"] = "=B1"; wb.save(path); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 200, response.text
    output = tmp_path / "formula.xlsx"; output.write_bytes(response.content)
    wb = openpyxl.load_workbook(output, data_only=False); assert wb["Main"]["C1"].value == "=B1"; wb.close()

    paths = make_period_books(tmp_path / "external", "quarterly")
    wb = openpyxl.load_workbook(paths[0]); wb["Main"]["C1"] = '=WEBSERVICE("https://example.invalid")'; wb.save(paths[0]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert "unsafe formula" in response.json()["detail"]


def test_external_relationship_is_rejected(tmp_path):
    paths = make_period_books(tmp_path, "quarterly")
    wb = openpyxl.load_workbook(paths[0]); wb["Main"]["C1"] = "External"; wb["Main"]["C1"].hyperlink = "https://example.invalid"; wb.save(paths[0]); wb.close()
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert "external workbook relationships" in response.json()["detail"]


def test_manifest_uses_one_read_only_open_per_source_and_closes_files(tmp_path, monkeypatch):
    paths = make_period_books(
        tmp_path, "quarterly",
        sheets_factory=lambda index: {
            "B2B": ("B2B", [row_for("B2B", index + 1)]),
            "CDNR": ("CDNR", [row_for("CDNR", index + 1)]),
        },
    )
    mapping = {f"s{index}": str(path) for index, path in enumerate(paths)}
    metadata = WorkbookInspector(mapping).inspect_all()
    calls = []
    original = source_manifest_module.openpyxl.load_workbook

    def tracked_load(path, *args, **kwargs):
        calls.append((str(path), kwargs.copy()))
        return original(path, *args, **kwargs)

    monkeypatch.setattr(source_manifest_module.openpyxl, "load_workbook", tracked_load)
    manifest = SourceManifestBuilder(mapping, metadata).build()
    assert len(calls) == len(paths)
    assert all(options.get("read_only") is True for _, options in calls)
    assert {key: value["row_count"] for key, value in manifest["categories"].items()} == {
        "b2b": 4, "cdnr": 4
    }
    assert len(list(iter_spooled_rows(manifest["categories"]["b2b"]))) == 4

    # Renaming succeeds on Windows only when every workbook/archive handle closed.
    for path in paths:
        moved = path.with_suffix(".closed")
        path.rename(moved)
        moved.rename(path)


def test_temporary_files_are_cleaned_after_failure(tmp_path, monkeypatch):
    paths = make_period_books(tmp_path / "inputs", "quarterly")
    wb = openpyxl.load_workbook(paths[0])
    wb["B2B"]["D2"] = "not-a-number"
    wb.save(paths[0])
    wb.close()
    runtime_temp = tmp_path / "runtime"
    runtime_temp.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(runtime_temp))
    response = run_core(paths, "quarterly", tmp_path)
    assert response.status_code == 400
    assert not list(runtime_temp.glob("gstr-merge-*"))
    assert not list(runtime_temp.glob("gstr-result-*"))
    assert not list((tmp_path / "inputs").glob("gstr-*.rows"))


def test_package_safety_accepts_a_valid_xlsx(tmp_path):
    path = make_book(tmp_path / "valid.xlsx")
    validate_xlsx_package(path, path.name)


def test_package_safety_rejects_non_ooxml_zip(tmp_path):
    path = tmp_path / "not-ooxml.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("document.txt", "not a workbook")
    with pytest.raises(PackageSafetyError, match="not a valid .xlsx"):
        validate_xlsx_package(path, path.name)


def test_package_safety_enforces_member_and_expanded_size_limits(tmp_path):
    path = make_book(tmp_path / "limited.xlsx")
    with zipfile.ZipFile(path) as archive:
        member_count = len(archive.infolist())
    with pytest.raises(PackageSafetyError, match="too many package entries") as members:
        validate_xlsx_package(path, path.name, max_members=member_count - 1)
    assert members.value.limit_exceeded is True
    with pytest.raises(PackageSafetyError, match="expanded workbook") as expanded:
        validate_xlsx_package(path, path.name, max_uncompressed_bytes=1)
    assert expanded.value.limit_exceeded is True


def test_package_safety_rejects_suspicious_compression_ratio(tmp_path):
    path = tmp_path / "compressed.xlsx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("xl/workbook.xml", "workbook")
        archive.writestr("xl/worksheets/sheet1.xml", "0" * 10_000)
    with pytest.raises(PackageSafetyError, match="unsafe workbook compression ratio"):
        validate_xlsx_package(path, path.name, max_compression_ratio=2)


def test_package_safety_rejects_detectable_encryption(monkeypatch, tmp_path):
    class Member:
        filename = "[Content_Types].xml"
        flag_bits = 0x1
        file_size = 1
        compress_size = 1

    class WorkbookMember(Member):
        filename = "xl/workbook.xml"
        flag_bits = 0

    class FakeArchive:
        def __enter__(self):
            return self

        def __exit__(self, *_):
            return False

        def infolist(self):
            return [Member(), WorkbookMember()]

    monkeypatch.setattr(package_safety_module.zipfile, "ZipFile", lambda _: FakeArchive())
    with pytest.raises(PackageSafetyError, match="password-protected"):
        validate_xlsx_package(tmp_path / "encrypted.xlsx", "encrypted.xlsx")


def test_inspector_enforces_configured_worksheet_dimensions(tmp_path):
    path = make_book(tmp_path / "dimensions.xlsx")
    errors = WorkbookInspector(
        {"one": str(path)}, max_rows=1, max_columns=3
    ).inspect_all()["one"]["errors"]
    assert any("exceeds safe dimensions" in error for error in errors)


@pytest.mark.parametrize(
    "context_overrides, expected",
    [
        ({"expected_gstin": "27AAAAA0000A1Z5"}, "does not match the expected GSTIN"),
        ({"expected_financial_year": "2026-27"}, "does not match the expected financial year"),
        ({"expected_periods": {"file_01": "Q2"}}, "supplied slot expects Q2"),
    ],
)
def test_processor_optionally_validates_expected_context(tmp_path, context_overrides, expected):
    paths = make_period_books(tmp_path, "quarterly")
    response = run_core(paths, "quarterly", tmp_path, **context_overrides)
    assert response.status_code == 400
    assert expected in response.text


def test_processor_removes_output_after_integrity_failure(tmp_path, monkeypatch):
    paths = make_period_books(tmp_path / "inputs", "quarterly")
    mapping = {f"file_{index:02d}": str(path) for index, path in enumerate(paths, 1)}
    filenames = {source_id: Path(path).name for source_id, path in mapping.items()}
    output = tmp_path / "must-be-removed.xlsx"

    monkeypatch.setattr(
        processor_module.IntegrityChecker,
        "verify",
        lambda self: ["synthetic independent verification failure"],
    )
    with pytest.raises(ProcessingValidationError, match="Integrity check failed"):
        process_workbooks(mapping, filenames, MergeContext("quarterly"), output)

    assert not output.exists()
    assert not list((tmp_path / "inputs").glob("gstr-*.rows"))
