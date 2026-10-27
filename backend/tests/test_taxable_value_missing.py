import json
from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from backend.app.models.records import (
    InvoiceNumberInfo, ParsedInvoiceRecord, RecordSource, SheetCategory, VoucherType,
)
from backend.app.models.reconciliation import ReconciliationStatus
from backend.app.services.excel_exporter import export_to_excel
from backend.app.services.file_parser import parse_file
from backend.app.services.json_exporter import generate_json_report
from backend.app.services.notes_reconciliation_engine import NotesReconciliationSummary
from backend.app.services.reconciliation_engine import reconcile_purchases
from backend.app.services.session_store import decode_session_value, encode_session_value
from backend.app.services.summary_generator import generate_text_summary


def _record(*, source=RecordSource.TALLY, invoice="INV-1", taxable="100", present=True,
            gstin="27AAAAA0000A1Z5", row=2, category=SheetCategory.PURCHASE_REGISTER):
    return ParsedInvoiceRecord(
        source=source, source_sheet="Purchase Register" if source == RecordSource.TALLY else "B2B",
        source_row=row, sheet_category=category,
        supplier_name="Supplier", gstin_original=gstin, gstin_normalized=gstin,
        gstin_valid=True, invoice_number=InvoiceNumberInfo(original=invoice, normalized=invoice),
        invoice_date=date(2026, 4, 1), voucher_type=VoucherType.PURCHASE,
        taxable_value=Decimal(taxable), taxable_value_present=present,
        cgst=Decimal("9"), sgst=Decimal("9"),
    )


def _gst(**kwargs):
    return _record(source=RecordSource.GST, category=SheetCategory.B2B, **kwargs)


def _write_accounts_xlsx(path, taxable):
    wb = Workbook(); ws = wb.active; ws.title = "Purchase Register"
    ws.append(["Supplier", "GSTIN", "Invoice No", "Date", "Taxable Value", "CGST", "SGST"])
    ws.append(["Supplier", "27AAAAA0000A1Z5", "INV-1", date(2026, 4, 1), taxable, 9, 9])
    wb.save(path)


def test_excel_blank_none_and_whitespace_are_missing_but_zero_is_present(tmp_path):
    for index, value in enumerate((None, "", "   ")):
        path = tmp_path / f"blank-{index}.xlsx"; _write_accounts_xlsx(path, value)
        record = parse_file(str(path), source_hint=RecordSource.TALLY).records[0]
        assert record.taxable_value_present is False
        assert reconcile_purchases([record], []).results[0].status == ReconciliationStatus.TAXABLE_VALUE_MISSING
    for index, value in enumerate((0, 0.00, "0.00")):
        path = tmp_path / f"zero-{index}.xlsx"; _write_accounts_xlsx(path, value)
        record = parse_file(str(path), source_hint=RecordSource.TALLY).records[0]
        assert record.taxable_value_present is True
        assert record.taxable_value == Decimal("0")
        assert reconcile_purchases([record], []).results[0].status != ReconciliationStatus.TAXABLE_VALUE_MISSING


def test_json_missing_null_and_zero_preserve_presence(tmp_path):
    path = tmp_path / "accounts.json"
    path.write_text(json.dumps({"Purchase Register": [
        {"Supplier": "A", "Invoice No": "1"},
        {"Supplier": "B", "Invoice No": "2", "Taxable Value": None},
        {"Supplier": "C", "Invoice No": "3", "Taxable Value": 0},
    ]}), encoding="utf-8")
    records = parse_file(str(path)).records
    assert [r.taxable_value_present for r in records] == [False, False, True]


def test_missing_taxable_overrides_financial_status_and_safe_exact_match_is_consumed():
    summary = reconcile_purchases([_record(present=False)], [_gst()])
    result = summary.results[0]
    assert result.status == ReconciliationStatus.TAXABLE_VALUE_MISSING
    assert result.invoice_number_gst == "INV-1"
    assert summary.taxable_value_missing == 1
    assert summary.major_differences == summary.perfectly_matched == summary.minor_differences == 0
    assert summary.not_in_tally == 0


def test_missing_taxable_without_safe_identity_does_not_consume_ambiguous_candidates():
    tally = _record(invoice="INV-X", gstin="", present=False)
    candidates = [_gst(invoice="INV-X", gstin="27AAAAA0000A1Z5"), _gst(invoice="INV-X", gstin="29BBBBB0000B1Z5", row=3)]
    summary = reconcile_purchases([tally], candidates)
    assert summary.results[0].status == ReconciliationStatus.TAXABLE_VALUE_MISSING
    assert summary.not_in_tally == 2


def test_cancelled_accounts_precedes_missing_taxable():
    record = _record(present=False)
    record.raw_data = {"Account": "(C A N C E L L E D) Supplier", "Vch/Bill No": "INV-1"}
    summary = reconcile_purchases([record], [])
    assert summary.cancelled_accounts_vouchers == 1
    assert summary.taxable_value_missing == 0
    assert not summary.results


def test_date_mismatch_remains_cross_cutting_for_safe_missing_taxable_match():
    gst = _gst(); gst.invoice_date = date(2026, 4, 2)
    result = reconcile_purchases([_record(present=False)], [gst]).results[0]
    assert result.status == ReconciliationStatus.TAXABLE_VALUE_MISSING
    assert result.date_mismatch is True


def test_session_round_trip_preserves_missing_and_explicit_zero():
    missing = _record(present=False); zero = _record(invoice="INV-0", taxable="0", present=True)
    restored = decode_session_value(encode_session_value({"records": [missing, zero]}))
    assert [r.taxable_value_present for r in restored["records"]] == [False, True]
    assert restored["records"][1].taxable_value == Decimal("0")


def test_excel_and_json_exports_distinguish_missing_from_zero(tmp_path):
    summary = reconcile_purchases([_record(present=False)], [_gst()])
    notes = NotesReconciliationSummary()
    output = tmp_path / "report.xlsx"; export_to_excel(str(output), summary, notes, "2026-27")
    wb = load_workbook(output, data_only=False)
    assert wb.sheetnames[:3] == ["Reconciliation Summary", "Missing & Differences", "Taxable Value Missing"]
    assert wb["Taxable Value Missing"]["H2"].value == "—"
    report = json.loads(generate_json_report(summary, notes, "2026-27"))
    assert report["counts"]["taxable_value_missing"] == 1
    assert report["taxable_value_missing"][0]["accounts_taxable_value"] is None


def test_summary_and_primary_invariant_include_taxable_missing():
    summary = reconcile_purchases([_record(present=False)], [])
    assert "Taxable Value Missing: 1" in generate_text_summary(summary, NotesReconciliationSummary())
    total = (summary.perfectly_matched + summary.minor_differences + summary.major_differences
             + summary.invoice_mismatches + summary.missing_in_gstr + summary.needs_review
             + summary.potential_duplicates + summary.cancelled_accounts_vouchers
             + summary.taxable_value_missing)
    assert total == summary.total_tally_invoices
