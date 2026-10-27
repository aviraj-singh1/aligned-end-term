"""Real-PostgreSQL worker lifecycle, race, and synthetic Excel integration tests."""

from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
import hashlib
from pathlib import Path
import shutil
import threading
from uuid import uuid4

import openpyxl
import pytest
from sqlalchemy.exc import SQLAlchemyError
from sqlmodel import Session, select

from backend.app.auth.security import utcnow
from backend.app.db.models import (
    AuditEvent,
    Client,
    ClientGSTIN,
    GstrMergeRun,
    User,
    Workspace,
    WorkspaceMembership,
)
from backend.app.services import gstr_merge_lifecycle_service as lifecycle
from backend.app.services import gstr_merge_worker_service as worker
from backend.app.services import object_storage
from backend.tests.gstr_merge.helpers import GSTIN, make_period_books
from backend.tests.test_database_schema import schema_engine


pytestmark = pytest.mark.postgres


@pytest.fixture(autouse=True)
def local_storage(monkeypatch):
    root = Path(__file__).resolve().parents[2] / "scratch" / "p4"
    shutil.rmtree(root, ignore_errors=True)
    monkeypatch.setenv("OBJECT_STORAGE_BACKEND", "local")
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("GSTR_LOCAL_STORAGE_ROOT", str(root))
    object_storage.gstr_storage_settings.cache_clear()
    object_storage.artifact_delivery_settings.cache_clear()
    yield root
    shutil.rmtree(root, ignore_errors=True)
    object_storage.gstr_storage_settings.cache_clear()
    object_storage.artifact_delivery_settings.cache_clear()


@pytest.fixture(autouse=True)
def isolate_worker_queue(schema_engine):
    """Prevent intentionally unfinished race fixtures from leaking to later tests."""
    yield
    with Session(schema_engine) as db:
        active = db.exec(select(GstrMergeRun).where(
            GstrMergeRun.status.in_(["QUEUED", "PROCESSING"])
        )).all()
        for run in active:
            run.status = "DELETED"
            run.deleted_at = utcnow()
            run.lease_token = None
            run.lease_expires_at = None
            run.next_attempt_at = None
            run.version += 1
            db.add(run)
        db.commit()


def _graph(engine):
    with Session(engine, expire_on_commit=False) as db:
        user = User(
            name="Synthetic worker owner",
            email=f"worker-{uuid4()}@example.test",
            email_normalized="placeholder@example.test",
            email_verified_at=utcnow(),
        )
        user.email_normalized = user.email
        db.add(user)
        db.flush()
        workspace = Workspace(name=f"Worker test {uuid4()}", type="FIRM", created_by=user.id)
        db.add(workspace)
        db.flush()
        db.add(WorkspaceMembership(
            workspace_id=workspace.id,
            user_id=user.id,
            role="OWNER",
            joined_at=utcnow(),
        ))
        client = Client(workspace_id=workspace.id, name="Synthetic client", created_by=user.id)
        db.add(client)
        db.flush()
        gstin = ClientGSTIN(workspace_id=workspace.id, client_id=client.id, gstin=GSTIN)
        db.add(gstin)
        db.commit()
        return user, workspace, client, gstin


def _queued_run(engine, tmp_path, mode="QUARTERLY", *, gstin=GSTIN, corrupt=False):
    user, workspace, client, registration = _graph(engine)
    sources = make_period_books(tmp_path / uuid4().hex, mode.lower())
    if gstin != GSTIN:
        for source in sources:
            book = openpyxl.load_workbook(source)
            book["Main"]["B1"] = gstin
            book.save(source)
            book.close()
    if corrupt:
        sources[0].write_bytes(b"not-an-xlsx-package")
    periods = lifecycle._periods(mode)
    run = GstrMergeRun(
        workspace_id=workspace.id,
        client_id=client.id,
        client_gstin_id=registration.id,
        created_by=user.id,
        mode=mode,
        financial_year="2025-26",
        status="QUEUED",
        next_attempt_at=utcnow(),
    )
    files = []
    for period, source in zip(periods, sources):
        reference = object_storage.gstr_input_reference(
            workspace.id, client.id, run.id, period
        )
        target = Path(reference)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        size = target.stat().st_size
        files.append({
            "period_code": period,
            "filename": source.name,
            "object_ref": reference,
            "expected_size_bytes": size,
            "actual_size_bytes": size,
            "content_type": object_storage.XLSX_MEDIA_TYPE,
            "upload_state": "UPLOADED",
        })
    run.input_manifest_json = {
        "schema_version": 1,
        "queued_at": utcnow().isoformat(),
        "files": files,
    }
    run_id = run.id
    with Session(engine, expire_on_commit=False) as db:
        db.add(run)
        db.commit()
    return user, workspace, client, registration, run_id, files


def _due(engine, run_id):
    with Session(engine) as db:
        run = db.get(GstrMergeRun, run_id)
        run.next_attempt_at = utcnow() - timedelta(seconds=1)
        db.add(run)
        db.commit()


@pytest.mark.parametrize("mode,count", [("QUARTERLY", 4), ("MONTHLY", 12)])
def test_synthetic_end_to_end_processing_and_authorized_download(
    schema_engine, tmp_path, monkeypatch, mode, count
):
    user, workspace, client, _, run_id, files = _queued_run(
        schema_engine, tmp_path, mode
    )
    attempt = worker.claim_next_run(engine=schema_engine)
    assert attempt and attempt.run_id == run_id and attempt.attempt_number == 1
    result = worker.process_attempt(attempt, engine=schema_engine)
    assert result.status == "COMPLETED"
    assert all(not Path(item["object_ref"]).exists() for item in files)

    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, run_id)
        assert run.status == "COMPLETED"
        assert run.output_size_bytes > 0 and len(run.output_sha256) == 64
        assert run.result_summary_json["files_processed"] == count
        assert run.result_summary_json["integrity_passed"] is True
        assert run.expires_at > utcnow()
        output = Path(run.output_object_ref)
        assert output.exists() and output.stat().st_size == run.output_size_bytes
        assert hashlib.sha256(output.read_bytes()).hexdigest() == run.output_sha256
        assert all("object_ref" not in item for item in run.input_manifest_json["files"])
        assert all(len(item["sha256"]) == 64 for item in run.input_manifest_json["files"])
        assert db.exec(select(AuditEvent).where(
            AuditEvent.entity_id == run_id,
            AuditEvent.action == "GSTR_MERGE_COMPLETED",
        )).one()
    book = openpyxl.load_workbook(output, read_only=True)
    assert book.sheetnames
    book.close()
    monkeypatch.setattr(lifecycle, "get_database_engine", lambda: schema_engine)
    response = lifecycle.download_upload_result(user, workspace.id, client.id, run_id)
    assert Path(response.path) == output
    with Session(schema_engine) as db:
        completed_version = db.get(GstrMergeRun, run_id).version
    lifecycle.delete_upload_run(
        user, workspace.id, client.id, run_id, completed_version
    )
    assert not output.exists()


def test_two_claimers_only_one_owns_single_run(schema_engine, tmp_path):
    *_, run_id, _ = _queued_run(schema_engine, tmp_path)
    barrier = threading.Barrier(2)

    def claim():
        barrier.wait()
        return worker.claim_next_run(engine=schema_engine)

    with ThreadPoolExecutor(max_workers=2) as pool:
        attempts = list(pool.map(lambda _index: claim(), range(2)))
    owned = [attempt for attempt in attempts if attempt is not None]
    assert len(owned) == 1 and owned[0].run_id == run_id
    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, run_id)
        assert run.status == "PROCESSING" and run.attempt_count == 1


def test_skip_locked_does_not_wait_for_locked_only_run(schema_engine, tmp_path):
    *_, run_id, _ = _queued_run(schema_engine, tmp_path)
    with Session(schema_engine) as locker:
        locker.exec(
            select(GstrMergeRun).where(GstrMergeRun.id == run_id).with_for_update()
        ).one()
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(worker.claim_next_run, engine=schema_engine)
            assert future.result(timeout=5) is None


def test_renewal_preserves_version_and_extends_lease(schema_engine, tmp_path):
    _queued_run(schema_engine, tmp_path)
    now = utcnow()
    settings = worker.WorkerSettings(lease_seconds=60, renew_seconds=10)
    attempt = worker.claim_next_run(engine=schema_engine, settings=settings, now=now)
    old_version = attempt.version
    worker.renew_lease(
        attempt, engine=schema_engine, settings=settings, now=now + timedelta(seconds=30)
    )
    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, attempt.run_id)
        assert run.version == old_version
        assert run.lease_expires_at == now + timedelta(seconds=90)


def test_expired_lease_requeues_and_stale_owner_cannot_complete(schema_engine, tmp_path):
    _queued_run(schema_engine, tmp_path)
    now = utcnow()
    settings = worker.WorkerSettings(
        lease_seconds=30, renew_seconds=5, retry_backoff_seconds=1
    )
    stale = worker.claim_next_run(engine=schema_engine, settings=settings, now=now)
    assert worker.recover_expired_runs(
        engine=schema_engine, settings=settings, now=now + timedelta(seconds=31)
    ) == 1
    reference = object_storage.gstr_output_reference(
        stale.workspace_id, stale.client_id, stale.run_id, stale.engine_version
    )
    with pytest.raises(worker.LeaseLostError):
        worker.complete_attempt(
            stale,
            output_reference=reference,
            output_size=1,
            output_sha256="a" * 64,
            summary={},
            engine=schema_engine,
            settings=settings,
        )
    with Session(schema_engine) as db:
        assert db.get(GstrMergeRun, stale.run_id).status == "QUEUED"


@pytest.mark.parametrize(
    "kind,expected_code",
    [("corrupt", "CORRUPT_WORKBOOK"), ("gstin", "GSTIN_MISMATCH")],
)
def test_deterministic_input_failure_is_terminal(
    schema_engine, tmp_path, kind, expected_code
):
    args = {"corrupt": True} if kind == "corrupt" else {"gstin": "07AAAAA0000A1Z4"}
    *_, run_id, files = _queued_run(schema_engine, tmp_path, **args)
    attempt = worker.claim_next_run(engine=schema_engine)
    result = worker.process_attempt(attempt, engine=schema_engine)
    assert result.status == "FAILED"
    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, run_id)
        assert run.status == "FAILED" and run.attempt_count == 1
        assert run.error_code == expected_code
        assert run.lease_token is None
    assert all(not Path(item["object_ref"]).exists() for item in files)


def test_transient_storage_failure_requeues_without_deleting_inputs(
    schema_engine, tmp_path, monkeypatch
):
    *_, run_id, files = _queued_run(schema_engine, tmp_path)
    attempt = worker.claim_next_run(engine=schema_engine)
    monkeypatch.setattr(
        object_storage,
        "download_private_object",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            object_storage.ObjectStorageUnavailableError("temporary")
        ),
    )
    result = worker.process_attempt(attempt, engine=schema_engine)
    assert result.status == "QUEUED"
    assert all(Path(item["object_ref"]).exists() for item in files)
    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, run_id)
        assert run.attempt_count == 1 and run.next_attempt_at > utcnow()


def test_repeated_infrastructure_failure_exhausts_budget_and_cleans_inputs(
    schema_engine, tmp_path, monkeypatch
):
    *_, run_id, files = _queued_run(schema_engine, tmp_path)
    monkeypatch.setattr(
        object_storage,
        "download_private_object",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            object_storage.ObjectStorageUnavailableError("temporary")
        ),
    )
    first = worker.claim_next_run(engine=schema_engine)
    assert worker.process_attempt(first, engine=schema_engine).status == "QUEUED"
    _due(schema_engine, run_id)
    second = worker.claim_next_run(engine=schema_engine)
    assert second.lease_token != first.lease_token
    assert worker.process_attempt(second, engine=schema_engine).status == "FAILED"
    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, run_id)
        assert run.attempt_count == 2 and run.error_code == "INFRASTRUCTURE_UNAVAILABLE"
    assert all(not Path(item["object_ref"]).exists() for item in files)


def test_equal_size_input_change_is_detected_from_first_attempt_digest(
    schema_engine, tmp_path, monkeypatch
):
    *_, run_id, files = _queued_run(schema_engine, tmp_path)
    original_processor = worker.process_workbooks
    with monkeypatch.context() as scoped:
        scoped.setattr(
            worker,
            "process_workbooks",
            lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("temporary")),
        )
        first = worker.claim_next_run(engine=schema_engine)
        assert worker.process_attempt(first, engine=schema_engine).status == "QUEUED"
    target = Path(files[0]["object_ref"])
    data = bytearray(target.read_bytes())
    data[-1] ^= 1
    target.write_bytes(data)
    assert target.stat().st_size == files[0]["actual_size_bytes"]
    _due(schema_engine, run_id)
    second = worker.claim_next_run(engine=schema_engine)
    assert worker.process_attempt(second, engine=schema_engine).status == "FAILED"
    with Session(schema_engine) as db:
        assert db.get(GstrMergeRun, run_id).error_code == "INPUT_CHANGED"
    assert worker.process_workbooks is original_processor


def test_missing_input_is_terminal(schema_engine, tmp_path):
    *_, run_id, files = _queued_run(schema_engine, tmp_path)
    attempt = worker.claim_next_run(engine=schema_engine)
    Path(files[0]["object_ref"]).unlink()
    assert worker.process_attempt(attempt, engine=schema_engine).status == "FAILED"
    with Session(schema_engine) as db:
        assert db.get(GstrMergeRun, run_id).error_code == "INPUT_MISSING"


def test_publication_failure_removes_unpublished_output_and_requeues(
    schema_engine, tmp_path, monkeypatch
):
    *_, run_id, _ = _queued_run(schema_engine, tmp_path)
    attempt = worker.claim_next_run(engine=schema_engine)
    reference = object_storage.gstr_output_reference(
        attempt.workspace_id, attempt.client_id, attempt.run_id, attempt.engine_version
    )
    monkeypatch.setattr(object_storage, "gstr_output_reference", lambda *_args: reference)
    monkeypatch.setattr(
        worker,
        "complete_attempt",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(SQLAlchemyError("temporary")),
    )
    assert worker.process_attempt(attempt, engine=schema_engine).status == "QUEUED"
    assert not Path(reference).exists()
    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, run_id)
        assert run.status == "QUEUED" and run.output_object_ref is None


def test_deleted_during_processing_cannot_publish(
    schema_engine, tmp_path, monkeypatch
):
    user, workspace, client, _, _, _ = _queued_run(schema_engine, tmp_path)
    attempt = worker.claim_next_run(engine=schema_engine)
    reference = object_storage.gstr_output_reference(
        attempt.workspace_id, attempt.client_id, attempt.run_id, attempt.engine_version
    )
    original = worker.process_workbooks

    def delete_after_bounded_core(*args, **kwargs):
        summary = original(*args, **kwargs)
        monkeypatch.setattr(lifecycle, "get_database_engine", lambda: schema_engine)
        lifecycle.delete_upload_run(
            user, workspace.id, client.id, attempt.run_id, attempt.version
        )
        return summary

    monkeypatch.setattr(worker, "process_workbooks", delete_after_bounded_core)
    monkeypatch.setattr(object_storage, "gstr_output_reference", lambda *_args: reference)
    assert worker.process_attempt(attempt, engine=schema_engine).status == "LEASE_LOST"
    assert not Path(reference).exists()
    with Session(schema_engine) as db:
        run = db.get(GstrMergeRun, attempt.run_id)
        assert run.status == "DELETED" and run.output_object_ref is None


def test_attempt_temporary_directory_is_removed(schema_engine, tmp_path, monkeypatch):
    _queued_run(schema_engine, tmp_path, corrupt=True)
    attempt = worker.claim_next_run(engine=schema_engine)
    real_factory = worker.tempfile.TemporaryDirectory
    created = []

    class TrackedTemporaryDirectory:
        def __init__(self, *args, **kwargs):
            self.inner = real_factory(*args, **kwargs)

        def __enter__(self):
            value = self.inner.__enter__()
            created.append(Path(value))
            return value

        def __exit__(self, *args):
            return self.inner.__exit__(*args)

    monkeypatch.setattr(worker.tempfile, "TemporaryDirectory", TrackedTemporaryDirectory)
    assert worker.process_attempt(attempt, engine=schema_engine).status == "FAILED"
    assert created and all(not path.exists() for path in created)


def test_graceful_stop_claims_no_new_work(schema_engine, tmp_path):
    *_, run_id, _ = _queued_run(schema_engine, tmp_path)
    stop = threading.Event()
    stop.set()
    worker.run_worker_loop(stop, engine=schema_engine)
    with Session(schema_engine) as db:
        assert db.get(GstrMergeRun, run_id).status == "QUEUED"
