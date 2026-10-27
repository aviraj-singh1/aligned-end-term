import json
from datetime import date
from decimal import Decimal
from openpyxl import load_workbook

from backend.app.models.records import (
    FinancialParseIssue, ParsedInvoiceRecord, RecordSource, SheetCategory, VoucherType,
)
from backend.app.models.reconciliation import (
    AggregatedGSTInvoice, AuditTrail, MatchMethod, ReconciliationResult,
    ReconciliationStatus, ReconciliationSummary,
)
from backend.app.services.json_exporter import generate_json_report
from backend.app.services.excel_exporter import export_to_excel
from backend.app.services.cancellation_detector import build_cancelled_voucher
from backend.app.services.manual_match_service import _execute_manual_match_in_session
from backend.app.services.normalizer import normalize_invoice_number
from backend.app.services.notes_reconciliation_engine import (
    AggregatedGSTCreditNote, NotesReconciliationSummary,
)
from backend.app.services.reconciliation_orchestrator import (
    _assign_result_ids, _compact_authoritative_state,
)
from backend.app.services.session_store import decode_session_value, encode_session_value


def row(source, number, *, note=False, raw_data=None):
    record = ParsedInvoiceRecord(
        source=source,
        source_sheet="Debit Note Register" if note and source == RecordSource.TALLY else "b2b",
        source_row=2,
        sheet_category=(SheetCategory.DEBIT_NOTE_REGISTER if note and source == RecordSource.TALLY
                        else SheetCategory.CDNR if note else
                        SheetCategory.PURCHASE_REGISTER if source == RecordSource.TALLY
                        else SheetCategory.B2B),
        supplier_name="Supplier", gstin_original="06AAAAA0000A1Z5",
        gstin_normalized="06AAAAA0000A1Z5", gstin_valid=True,
        invoice_number=normalize_invoice_number(number), invoice_date=date(2026, 4, 1),
        voucher_type=(VoucherType.DEBIT_NOTE if note and source == RecordSource.TALLY
                      else VoucherType.CREDIT_NOTE if note else VoucherType.PURCHASE),
        taxable_value=Decimal("100.00"), cgst=Decimal("9.00"), sgst=Decimal("9.00"),
        raw_data=raw_data or {"Particulars": "Supplier", "Voucher No.": number,
                              "GSTIN/UIN": "06AAAAA0000A1Z5", "Value": 100},
    )
    if note:
        record.note_number = number
        record.note_date = date(2026, 4, 1)
        record.note_type = "C" if source == RecordSource.GST else ""
    return record


def purchase_review(*, missing=False, malformed=False):
    tally = row(RecordSource.TALLY, "T-1")
    tally.taxable_value_present = not missing
    if malformed:
        tally.financial_parse_issues = [FinancialParseIssue("cgst", "9OO", "invalid cgst")]
    gst_row = row(RecordSource.GST, "G-1")
    candidate = AggregatedGSTInvoice(
        gstin=gst_row.gstin_normalized, supplier_name=gst_row.supplier_name,
        invoice_number_original="G-1", invoice_number_normalized="G1",
        invoice_date=gst_row.invoice_date, taxable_value=gst_row.taxable_value,
        cgst=gst_row.cgst, sgst=gst_row.sgst, source_rows=[gst_row],
    )
    result = ReconciliationResult(
        supplier_gstin=tally.gstin_normalized, supplier_tally=tally.supplier_name,
        invoice_number_tally="T-1", taxable_value=tally.taxable_value,
        taxable_value_missing=missing,
        status=ReconciliationStatus.TAXABLE_VALUE_MISSING if missing else ReconciliationStatus.NEEDS_REVIEW,
        match_method=MatchMethod.AMBIGUOUS,
        audit=AuditTrail(tally_rows=[tally], candidates=[candidate], candidate_count=1),
    )
    summary = ReconciliationSummary(
        total_tally_invoices=1, total_gst_invoices=1, results=[result],
        not_in_tally_records=[candidate],
    )
    summary.compute_counts()
    return summary, NotesReconciliationSummary(), result, candidate


def test_compact_purchase_manual_match_survives_round_trip_and_consumes_once():
    purchases, notes, result, candidate = purchase_review()
    match_map = {}
    _assign_result_ids(purchases.results, match_map)
    _compact_authoritative_state(purchases, notes)
    restored = decode_session_value(encode_session_value({
        "reconciliation_summary": purchases, "notes_summary": notes,
        "manual_match_map": match_map, "reports": {},
    }))
    result_id = result.api_result_id
    candidate_id = candidate.api_candidate_id
    assert restored["manual_match_map"][result_id][candidate_id] is True
    response = _execute_manual_match_in_session(restored, result_id, candidate_id)
    assert response["new_status"] == ReconciliationStatus.PERFECTLY_MATCHED.value
    assert restored["reconciliation_summary"].not_in_tally == 0
    try:
        _execute_manual_match_in_session(restored, result_id, candidate_id)
        assert False, "candidate was consumed twice"
    except ValueError:
        pass


def test_compact_notes_and_taxable_missing_manual_matches_survive_round_trip():
    tally = row(RecordSource.TALLY, "DN-1", note=True)
    gst_row = row(RecordSource.GST, "CN-1", note=True)
    candidate = AggregatedGSTCreditNote(
        gstin=gst_row.gstin_normalized, supplier_name="Supplier",
        note_number_original="CN-1", note_number_normalized="CN1",
        note_date=gst_row.note_date, taxable_value=Decimal("100"),
        cgst=Decimal("9"), sgst=Decimal("9"), source_rows=[gst_row],
    )
    note_result = ReconciliationResult(
        supplier_gstin=tally.gstin_normalized, invoice_number_tally="DN-1",
        status=ReconciliationStatus.NEEDS_REVIEW, match_method=MatchMethod.AMBIGUOUS,
        audit=AuditTrail(tally_rows=[tally], candidates=[candidate]),
    )
    notes = NotesReconciliationSummary(
        total_tally_debit_notes=1, total_gst_credit_notes=1,
        results=[note_result], not_in_tally_notes=[candidate],
    )
    notes.compute_counts()
    purchases, _, missing_result, purchase_candidate = purchase_review(missing=True)
    match_map = {}
    _assign_result_ids(purchases.results, match_map)
    _assign_result_ids(notes.results, match_map)
    _compact_authoritative_state(purchases, notes)
    restored = decode_session_value(encode_session_value({
        "reconciliation_summary": purchases, "notes_summary": notes,
        "manual_match_map": match_map, "reports": {},
    }))
    note_response = _execute_manual_match_in_session(
        restored, note_result.api_result_id, candidate.api_candidate_id)
    missing_response = _execute_manual_match_in_session(
        restored, missing_result.api_result_id, purchase_candidate.api_candidate_id)
    assert note_response["new_status"] == ReconciliationStatus.PERFECTLY_MATCHED.value
    assert missing_response["new_status"] == ReconciliationStatus.TAXABLE_VALUE_MISSING.value


def test_compaction_preserves_malformed_amendment_and_cancellation_evidence():
    purchases, notes, result, candidate = purchase_review(malformed=True)
    cancelled_row = row(RecordSource.TALLY, "CANCEL-1", raw_data={"Remarks": "Cancelled"})
    purchases.cancelled_vouchers = [build_cancelled_voucher(cancelled_row)]
    candidate.source_rows[0].sheet_category = SheetCategory.B2BA
    candidate.source_rows[0].amendment_original_number = "OLD-1"
    candidate.source_rows[0].amendment_original_date = date(2026, 3, 1)
    _compact_authoritative_state(purchases, notes)
    restored = decode_session_value(encode_session_value(purchases))
    tally = restored.results[0].audit.tally_rows[0]
    amended = restored.not_in_tally_records[0].source_rows[0]
    assert tally.raw_data == {}
    assert tally.financial_parse_issues[0].raw_value == "9OO"
    assert amended.sheet_category == SheetCategory.B2BA
    assert amended.amendment_original_number == "OLD-1"
    assert restored.cancelled_vouchers[0].cancellation_evidence == "Remarks: Cancelled"

    match_map = {}
    _assign_result_ids(restored.results, match_map)
    session = {"reconciliation_summary": restored, "notes_summary": notes,
               "manual_match_map": match_map, "reports": {}}
    try:
        _execute_manual_match_in_session(
            session, restored.results[0].api_result_id,
            restored.results[0].audit.candidates[0].api_candidate_id,
        )
        assert False, "malformed record became manually matchable"
    except ValueError as error:
        assert "invalid monetary" in str(error)


def test_compact_export_after_round_trip_keeps_source_system_and_full_json(tmp_path):
    purchases, notes, _, _ = purchase_review()
    purchases.results[0].audit.tally_rows[0].raw_data = {
        "Account": "Supplier", "Vch/Bill No": "T-1",
        "TIN/GSTIN No.": "06AAAAA0000A1Z5", "Taxable Amt.": 100,
    }
    _compact_authoritative_state(purchases, notes)
    restored = decode_session_value(encode_session_value({"p": purchases, "n": notes}))
    report = json.loads(generate_json_report(restored["p"], restored["n"], "Apr-26"))
    assert report["reconciliation_summary"][0]["source_system"] == "BUSY"
    assert "not_in_tally" in report and "notes" in report
    target = tmp_path / "compact.xlsx"
    export_to_excel(str(target), restored["p"], restored["n"], "Apr-26")
    workbook = load_workbook(target, read_only=True)
    assert workbook.sheetnames == [
        "Reconciliation Summary", "Missing & Differences", "Taxable Value Missing",
        "Date Mismatches", "Invoice Mismatches", "Cancelled Vouchers",
        "Invoice No Remarks", "Not in Tally", "Notes",
    ]


def test_compacted_authoritative_state_is_materially_smaller():
    purchases = ReconciliationSummary()
    for index in range(100):
        tally = row(RecordSource.TALLY, f"T-{index}", raw_data={
            f"Column {column}": f"large repeated source value {index}-{column}"
            for column in range(20)
        })
        result = ReconciliationResult(
            invoice_number_tally=f"T-{index}", status=ReconciliationStatus.MISSING_IN_GSTR,
            audit=AuditTrail(tally_rows=[tally]),
        )
        purchases.results.append(result)
    notes = NotesReconciliationSummary()
    before = len(json.dumps(encode_session_value({"p": purchases, "n": notes})))
    _compact_authoritative_state(purchases, notes)
    after = len(json.dumps(encode_session_value({"p": purchases, "n": notes})))
    assert after < before
    assert before - after > before * 0.40
