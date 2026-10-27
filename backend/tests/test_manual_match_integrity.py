from datetime import date
from decimal import Decimal
import json
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.main import app
from backend.app.models.records import ParsedInvoiceRecord, RecordSource
from backend.app.models.reconciliation import (
    AggregatedGSTInvoice,
    AuditTrail,
    MatchMethod,
    ReconciliationResult,
    ReconciliationStatus,
    ReconciliationSummary,
)
from backend.app.services.normalizer import normalize_invoice_number
from backend.app.services.notes_reconciliation_engine import (
    AggregatedGSTCreditNote,
    NotesReconciliationSummary,
)
from backend.app.services.excel_exporter import export_to_excel
from backend.app.services.json_exporter import generate_json_report
from backend.app.services.session_store import decode_session_value, encode_session_value
from backend.app.services.reconciliation_orchestrator import _build_reconciliation_response
from backend.app.services.session_manager import create_session, get_session, update_session


client = TestClient(app)


def _manual_session(
    *,
    tally_date=date(2026, 3, 1),
    gst_date=date(2026, 3, 1),
    tally_taxable="100.00",
    gst_taxable="100.00",
):
    tally = ParsedInvoiceRecord(
        source=RecordSource.TALLY,
        supplier_name="Tally Supplier",
        gstin_original="06AAAAA0000A1Z5",
        gstin_normalized="06AAAAA0000A1Z5",
        invoice_number=normalize_invoice_number("TALLY-001"),
        invoice_date=tally_date,
        taxable_value=Decimal(tally_taxable),
        cgst=Decimal("9.00"),
        sgst=Decimal("9.00"),
    )
    gst_row = ParsedInvoiceRecord(
        source=RecordSource.GST,
        supplier_name="GST Supplier",
        gstin_normalized="06AAAAA0000A1Z5",
        invoice_number=normalize_invoice_number("GSTR-001"),
        invoice_date=gst_date,
        taxable_value=Decimal(gst_taxable),
        cgst=Decimal("9.00"),
        sgst=Decimal("9.00"),
    )
    candidate = AggregatedGSTInvoice(
        gstin=gst_row.gstin_normalized,
        supplier_name=gst_row.supplier_name,
        invoice_number_original=gst_row.invoice_number.original,
        invoice_number_normalized=gst_row.invoice_number.normalized,
        invoice_date=gst_date,
        taxable_value=gst_row.taxable_value,
        cgst=gst_row.cgst,
        sgst=gst_row.sgst,
        source_rows=[gst_row],
    )
    result_id = str(uuid.uuid4())
    candidate_id = str(uuid.uuid4())
    candidate.api_candidate_id = candidate_id
    review = ReconciliationResult(
        supplier_gstin=tally.gstin_normalized,
        supplier_tally=tally.supplier_name,
        invoice_number_tally=tally.invoice_number.original,
        invoice_date=tally_date,
        taxable_value=tally.taxable_value,
        status=ReconciliationStatus.NEEDS_REVIEW,
        match_method=MatchMethod.AMBIGUOUS,
        audit=AuditTrail(
            tally_rows=[tally],
            tally_taxable=tally.taxable_value,
            candidates=[candidate],
            candidate_count=1,
        ),
    )
    review.api_result_id = result_id
    not_in_tally = ReconciliationResult(
        supplier_gstin=candidate.gstin,
        supplier_gst=candidate.supplier_name,
        invoice_number_gst=candidate.invoice_number_original,
        invoice_date=gst_date,
        taxable_value=candidate.taxable_value,
        status=ReconciliationStatus.NOT_IN_TALLY,
        match_method=MatchMethod.NO_MATCH,
        audit=AuditTrail(
            gst_rows=candidate.source_rows,
            gst_invoice_original=candidate.invoice_number_original,
            gst_invoice_normalized=candidate.invoice_number_normalized,
            gst_gstin=candidate.gstin,
            gst_supplier=candidate.supplier_name,
            gst_taxable=candidate.taxable_value,
        ),
    )
    summary = ReconciliationSummary(
        total_tally_invoices=1,
        total_gst_invoices=1,
        results=[review, not_in_tally],
        not_in_tally_records=[candidate],
    )
    summary.compute_counts()
    session_id = create_session()
    update_session(session_id, "reconciliation_summary", summary)
    update_session(session_id, "notes_summary", NotesReconciliationSummary())
    update_session(session_id, "manual_match_map", {result_id: {candidate_id: candidate}})
    return session_id, result_id, candidate_id, candidate


def _notes_manual_session():
    tally = ParsedInvoiceRecord(
        source=RecordSource.TALLY,
        supplier_name="Accounting Supplier",
        gstin_normalized="06AAAAA0000A1Z5",
        invoice_number=normalize_invoice_number("DN-001"),
        invoice_date=date(2026, 3, 1),
        taxable_value=Decimal("100.00"),
        cgst=Decimal("9.00"),
        sgst=Decimal("9.00"),
    )
    gst_row = ParsedInvoiceRecord(
        source=RecordSource.GST,
        supplier_name="GST Supplier",
        gstin_normalized="06AAAAA0000A1Z5",
        invoice_number=normalize_invoice_number("CN-002"),
        note_number="CN-002",
        note_date=date(2026, 3, 1),
        taxable_value=Decimal("100.00"),
        cgst=Decimal("9.00"),
        sgst=Decimal("9.00"),
    )
    candidate = AggregatedGSTCreditNote(
        gstin=gst_row.gstin_normalized,
        supplier_name=gst_row.supplier_name,
        note_number_original="CN-002",
        note_number_normalized="CN002",
        note_date=gst_row.note_date,
        taxable_value=gst_row.taxable_value,
        cgst=gst_row.cgst,
        sgst=gst_row.sgst,
        source_rows=[gst_row],
    )
    result_id = str(uuid.uuid4())
    candidate_id = str(uuid.uuid4())
    candidate.api_candidate_id = candidate_id
    review = ReconciliationResult(
        supplier_gstin=tally.gstin_normalized,
        supplier_tally=tally.supplier_name,
        invoice_number_tally=tally.invoice_number.original,
        invoice_date=tally.invoice_date,
        taxable_value=tally.taxable_value,
        status=ReconciliationStatus.NEEDS_REVIEW,
        match_method=MatchMethod.AMBIGUOUS,
        audit=AuditTrail(
            tally_rows=[tally], tally_taxable=tally.taxable_value,
            candidates=[candidate], candidate_count=1,
        ),
    )
    review.api_result_id = result_id
    synthetic = ReconciliationResult(
        supplier_gstin=candidate.gstin,
        supplier_gst=candidate.supplier_name,
        invoice_number_gst=candidate.note_number_original,
        invoice_date=candidate.note_date,
        taxable_value=candidate.taxable_value,
        status=ReconciliationStatus.NOT_IN_TALLY,
        match_method=MatchMethod.NO_MATCH,
        audit=AuditTrail(
            gst_rows=candidate.source_rows,
            gst_invoice_original=candidate.note_number_original,
            gst_invoice_normalized=candidate.note_number_normalized,
            gst_gstin=candidate.gstin,
            gst_supplier=candidate.supplier_name,
            gst_taxable=candidate.taxable_value,
        ),
    )
    notes = NotesReconciliationSummary(
        total_tally_debit_notes=1,
        total_gst_credit_notes=1,
        results=[review, synthetic],
        not_in_tally_notes=[candidate],
    )
    notes.compute_counts()
    session_id = create_session()
    update_session(session_id, "reconciliation_summary", ReconciliationSummary())
    update_session(session_id, "notes_summary", notes)
    update_session(session_id, "manual_match_map", {result_id: {candidate_id: candidate}})
    return session_id, result_id, candidate_id, candidate


def test_review_payload_uses_raw_tally_values_and_candidate_date():
    session_id, _, _, _ = _manual_session()
    session = get_session(session_id)
    review = session["reconciliation_summary"].results[0]
    review.audit.igst_comparison.tally_value = Decimal("0.00")
    review.audit.cgst_comparison.tally_value = Decimal("0.00")
    review.audit.sgst_comparison.tally_value = Decimal("0.00")
    data = _build_reconciliation_response(
        session_id, session["reconciliation_summary"], session["notes_summary"]
    )
    payload = next(r for r in data["results"] if r["accounts_invoice"] == "TALLY-001")
    assert payload["accounts_supplier"] == "Tally Supplier"
    assert payload["accounts_gstin"] == "06AAAAA0000A1Z5"
    assert payload["accounts_taxable"] == "100.00"
    assert payload["accounts_igst"] == "0.00"
    assert payload["accounts_cgst"] == "9.00"
    assert payload["accounts_sgst"] == "9.00"
    assert payload["accounts_date"] == "2026-03-01"
    assert payload["candidates"][0]["invoice_date"] == "2026-03-01"


def test_review_payload_preserves_unavailable_tally_source_as_null():
    session_id, _, _, _ = _manual_session()
    session = get_session(session_id)
    review = session["reconciliation_summary"].results[0]
    review.audit.tally_rows = []
    data = _build_reconciliation_response(
        session_id, session["reconciliation_summary"], session["notes_summary"]
    )
    payload = next(r for r in data["results"] if r["accounts_invoice"] == "TALLY-001")
    assert payload["accounts_taxable"] is None
    assert payload["accounts_igst"] is None
    assert payload["accounts_cgst"] is None
    assert payload["accounts_sgst"] is None
    assert payload["accounts_date"] is None


def _match(session_id, result_id, candidate_id):
    return client.post(
        f"/api/manual-match?session_id={session_id}",
        json={"result_id": result_id, "candidate_id": candidate_id},
    )


def _result_items(session_id, category="all"):
    response = client.get(f"/api/session/{session_id}/results", params={"category": category, "page_size": 100})
    assert response.status_code == 200
    return response.json()["items"]


@pytest.mark.parametrize(
    ("gst_taxable", "expected"),
    [("100.00", "Perfectly Matched"), ("100.50", "Minor Rounding Difference"), ("102.00", "Major Difference")],
)
def test_manual_match_financial_classification(gst_taxable, expected):
    session_id, result_id, candidate_id, _ = _manual_session(gst_taxable=gst_taxable)
    response = _match(session_id, result_id, candidate_id)
    assert response.status_code == 200
    matched = next(r for r in _result_items(session_id) if r["accounts_invoice"] == "TALLY-001")
    assert matched["status"] == expected
    assert matched["match_method"] == "Manual Match"


@pytest.mark.parametrize(
    ("tally_date", "gst_date", "mismatch", "days"),
    [
        (date(2026, 3, 1), date(2026, 3, 1), False, 0),
        (date(2026, 3, 1), date(2026, 2, 28), True, 1),
        (date(2026, 3, 10), date(2026, 2, 27), True, 11),
        (date(2026, 3, 1), None, False, None),
    ],
)
def test_manual_match_date_metadata(tally_date, gst_date, mismatch, days):
    session_id, result_id, candidate_id, _ = _manual_session(tally_date=tally_date, gst_date=gst_date)
    data = _match(session_id, result_id, candidate_id).json()
    matched = next(r for r in _result_items(session_id) if r["accounts_invoice"] == "TALLY-001")
    assert matched["date_mismatch"] is mismatch
    assert matched["date_difference_days"] == days
    assert data["counts"]["date_mismatches"] == int(mismatch)


def test_manual_match_returns_authoritative_reconciliation_response():
    session_id, result_id, candidate_id, _ = _manual_session()
    data = _match(session_id, result_id, candidate_id).json()
    assert data["session_id"] == session_id
    assert set(("counts", "results", "not_in_tally", "notes", "summary_text")) <= data.keys()


def test_manual_match_consumes_once_and_removes_not_in_tally():
    session_id, result_id, candidate_id, candidate = _manual_session()
    data = _match(session_id, result_id, candidate_id).json()
    assert candidate.consumed is True
    assert data["not_in_tally"] == []
    assert data["counts"]["not_in_tally"] == 0
    assert all(r["status"] != "Not in Tally" for r in data["results"])


def test_manual_match_recomputes_primary_partition():
    session_id, result_id, candidate_id, _ = _manual_session(gst_taxable="102.00")
    counts = _match(session_id, result_id, candidate_id).json()["counts"]
    assert counts["major_difference"] == 1
    assert counts["needs_review_ambiguous"] == 0
    assert sum(counts[k] for k in (
        "perfectly_matched", "minor_rounding_difference", "major_difference",
        "invoice_mismatches", "missing_in_gstr_2b", "needs_review_ambiguous",
    )) == counts["accounts_invoices"]


def test_second_candidate_consumption_returns_409():
    session_id, result_id, candidate_id, _ = _manual_session()
    assert _match(session_id, result_id, candidate_id).status_code == 200
    assert _match(session_id, result_id, candidate_id).status_code == 409


def test_notes_manual_match_removes_unmatched_state_and_updates_counts():
    session_id, result_id, candidate_id, candidate = _notes_manual_session()
    response = _match(session_id, result_id, candidate_id)
    assert response.status_code == 200
    session = get_session(session_id)
    notes = session["notes_summary"]
    assert candidate.consumed is True
    assert notes.not_in_tally_notes == []
    assert not any(result.status == ReconciliationStatus.NOT_IN_TALLY for result in notes.results)
    assert notes.not_in_tally == 0
    assert notes.needs_review == 0
    assert notes.perfectly_matched == 1
    page = _result_items(session_id, "notes")
    assert len(page) == 1
    assert page[0]["accounts_invoice"] == "DN-001"
    assert page[0]["gst_invoice"] == "CN-002"


def test_notes_candidate_cannot_be_consumed_twice():
    session_id, result_id, candidate_id, _ = _notes_manual_session()
    assert _match(session_id, result_id, candidate_id).status_code == 200
    assert _match(session_id, result_id, candidate_id).status_code == 409


def test_notes_manual_match_state_survives_serialization_reload():
    session_id, result_id, candidate_id, _ = _notes_manual_session()
    assert _match(session_id, result_id, candidate_id).status_code == 200
    restored = decode_session_value(encode_session_value(get_session(session_id)))
    notes = restored["notes_summary"]
    assert notes.not_in_tally_notes == []
    assert notes.not_in_tally == 0
    assert len(notes.results) == 1
    assert notes.results[0].match_method == MatchMethod.MANUAL_MATCH
    assert restored["manual_match_map"][result_id][candidate_id].consumed is True


def test_notes_manual_match_normalizes_opposite_signs_only_in_notes_domain():
    session_id, result_id, candidate_id, candidate = _notes_manual_session()
    session = get_session(session_id)
    note_result = session["notes_summary"].results[0]
    tally = note_result.audit.tally_rows[0]
    tally.taxable_value = Decimal("-100.00")
    tally.cgst = tally.sgst = Decimal("-9.00")
    candidate.taxable_value = Decimal("100.00")
    candidate.cgst = candidate.sgst = Decimal("9.00")
    response = _match(session_id, result_id, candidate_id)
    assert response.status_code == 200
    item = _result_items(session_id, "notes")[0]
    assert item["status"] == "Perfectly Matched"
    assert item["accounts_taxable"] == "-100.00"
    assert item["gst_taxable"] == "100.00"
    assert item["taxable_difference"] == "0.00"


def test_manual_match_rejects_cross_domain_candidate_even_if_map_is_tampered():
    session_id, result_id, candidate_id, _ = _manual_session()
    session = get_session(session_id)
    note_candidate = AggregatedGSTCreditNote(
        gstin="06AAAAA0000A1Z5", note_number_original="CN-X",
        note_number_normalized="CNX", taxable_value=Decimal("100.00"),
    )
    note_candidate.api_candidate_id = candidate_id
    session["manual_match_map"] = {result_id: {candidate_id: note_candidate}}
    response = _match(session_id, result_id, candidate_id)
    assert response.status_code == 400
    assert "purchase candidates" in response.json()["detail"].lower()


def test_excel_after_notes_manual_match_has_only_the_matched_note(tmp_path):
    session_id, result_id, candidate_id, _ = _notes_manual_session()
    assert _match(session_id, result_id, candidate_id).status_code == 200
    session = get_session(session_id)
    path = tmp_path / "notes-manual-match.xlsx"
    export_to_excel(
        str(path), session["reconciliation_summary"], session["notes_summary"], "March 2026"
    )
    rows = list(load_workbook(path)["Notes"].iter_rows(min_row=2, values_only=True))
    assert len([row for row in rows if row[5] == "CN-002"]) == 1
    assert not any(row[14] == ReconciliationStatus.NOT_IN_TALLY.value for row in rows)


def test_json_after_notes_manual_match_has_only_the_matched_note():
    session_id, result_id, candidate_id, _ = _notes_manual_session()
    assert _match(session_id, result_id, candidate_id).status_code == 200
    session = get_session(session_id)
    data = json.loads(generate_json_report(
        session["reconciliation_summary"], session["notes_summary"], "March 2026"
    ))
    exported = [item for item in data["notes"] if item.get("invoice_number_gst") == "CN-002"]
    assert len(exported) == 1
    assert exported[0]["status"] == ReconciliationStatus.PERFECTLY_MATCHED.value


def test_purchase_manual_match_contract_remains_unchanged():
    session_id, result_id, candidate_id, _ = _manual_session()
    assert _match(session_id, result_id, candidate_id).status_code == 200
    session = get_session(session_id)
    purchases = session["reconciliation_summary"]
    assert purchases.not_in_tally_records == []
    assert purchases.not_in_tally == 0
    assert len(purchases.results) == 1
    assert purchases.results[0].match_method == MatchMethod.MANUAL_MATCH


def test_unknown_candidate_returns_400():
    session_id, result_id, _, _ = _manual_session()
    assert _match(session_id, result_id, str(uuid.uuid4())).status_code == 400


def test_unknown_result_returns_400():
    session_id, _, candidate_id, _ = _manual_session()
    assert _match(session_id, str(uuid.uuid4()), candidate_id).status_code == 400


def test_unknown_session_returns_404():
    assert _match(str(uuid.uuid4()), str(uuid.uuid4()), str(uuid.uuid4())).status_code == 404


def test_session_get_remains_lightweight_status_contract():
    session_id, _, _, _ = _manual_session()
    response = client.get(f"/api/session/{session_id}")
    assert response.status_code == 200
    assert set(response.json()) == {
        "session_id", "created_at", "last_accessed", "has_inspection", "has_reconciliation"
    }
